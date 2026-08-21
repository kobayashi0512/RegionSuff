"""Insert the RegionSuff GitHub archive link into the manuscript availability statement."""

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_3.3.docx')
OUTPUT = Path(__file__).resolve().parents[1] / 'manuscript' / 'RegionSuff_3.3_GitHubArchive.docx'
URL = 'https://github.com/kobayashi0512/RegionSuff'
OLD = ('The analysis code, model configurations, frozen experiment outputs, paired-bootstrap results, '
       'validation-list audit, manuscript builder, and figure source files are available from the '
       'corresponding author upon reasonable request.')
PREFIX = ('The analysis code, model configurations, frozen experiment outputs, paired-bootstrap results, '
          'validation-list audit, manuscript builder, and editable figure source files are archived at ')
SUFFIX = ('. Dataset and model reuse remains subject to the corresponding source licenses and terms.')


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    r_pr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    r_pr.append(underline)
    run.append(r_pr)
    content = OxmlElement('w:t')
    content.text = text
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def main():
    doc = Document(SOURCE)
    matches = [p for p in doc.paragraphs if OLD in p.text]
    if len(matches) != 1:
        raise RuntimeError(f'Expected exactly one availability paragraph, found {len(matches)}')
    p = matches[0]
    target_runs = [run for run in p.runs if OLD in run.text]
    if len(target_runs) != 1:
        raise RuntimeError(f'Expected exactly one target run, found {len(target_runs)}')
    target = target_runs[0]
    target.text = target.text.replace(OLD, PREFIX)
    add_hyperlink(p, URL, URL)
    p.add_run(SUFFIX)
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip() == 'Funding':
            remove_paragraph(paragraph)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f'Wrote {OUTPUT}')


if __name__ == '__main__':
    main()
