from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path("/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.3.docx")
OUTPUT = Path("/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.3.updated.docx")
FIGURE = Path(
    "/Users/tongyue/Documents/77f /7hao/paper_design/ppt_build/"
    "figure1_large_tnr/final-render/slide-1.png"
)


ABSTRACT = [
    (
        "Reliable AI-assisted retinal diagnosis depends not only on image clarity but also on "
        "whether the acquired field contains the evidence required for the target task. Existing "
        "retinal image-quality assessment methods mainly evaluate blur, illumination, contrast, "
        "and imaging artifacts. A technically gradable image may nevertheless omit macular, "
        "optic-disc, or peripheral signs relevant to diabetic retinopathy (DR) grading, leaving "
        "diagnostic sufficiency unexamined."
    ),
    (
        "We introduce RegionSuff, a visibility-conditioned model for evaluating DR classification "
        "under restricted retinal coverage in ultra-widefield fundus imaging. RegionSuff divides "
        "the retinal foreground into a 3×3 evidence bank, extracts regional features using a frozen "
        "retinal image-quality encoder, and applies masked attention to aggregate only visible "
        "regions. Random regional masking introduces controlled evidence loss during training, "
        "while regional masks and pixel apertures examine how evidence composition, area, location, "
        "shape, and topology affect prediction. The model was developed using 1,630 UWF-DR records "
        "and evaluated without target-domain adaptation or model selection on the official "
        "2,597-image MMRDR-UWF test set."
    ),
    (
        "On UWF-DR, macro-F1 decreased from 0.830 with the full view to 0.790 with the center-only "
        "view and increased to 0.829 when the four axial neighboring regions were restored. On "
        "MMRDR-UWF, macro-F1 was 0.625, 0.600, and 0.642 for full-view, center-only, and "
        "center-plus-cross inference, respectively. Equal-area apertures with different locations "
        "or shapes produced different outcomes, and central-region occlusion caused the largest "
        "performance loss. A paired-view model detected limited-view underestimation with an AUROC "
        "of 0.740 on the held-out split, although grouped nested evaluation reduced the AUROC to "
        "0.653. RegionSuff extends retinal image-quality assessment from photographic quality to "
        "the availability and spatial organization of diagnostic evidence."
    ),
]

KEYWORDS = (
    "diagnostic sufficiency; retinal image-quality assessment; ultra-widefield fundus imaging; "
    "diabetic retinopathy; masked attention; limited-view underestimation"
)

STUDY_OVERVIEW = [
    (
        "Using publicly available, de-identified data, we conducted a retrospective computational "
        "study without additional image acquisition. Model performance was evaluated on the held-out "
        "test set under three predefined visibility settings: full-view, center-only, and "
        "center-plus-cross. Macro-F1 for three-class DR grading was the primary outcome, with accuracy, "
        "ordinal severity MAE, and underestimation rate included as secondary outcomes. Source-domain "
        "training was repeated across five random seeds to capture variability arising from model "
        "optimization. For analyses with per-image predictions, differences between visibility "
        "settings were summarized using image-level paired bootstrap 95% confidence intervals. [28]"
    ),
    (
        "The analyses were designed to characterize how the evidence available to the DR classifier "
        "changed as retinal coverage was restricted. The masks therefore represent controlled, "
        "synthetic restrictions of the observed retina rather than real pairs of conventional-field "
        "and UWF acquisitions. Because the datasets do not provide clinician-adjudicated labels of "
        "diagnostic sufficiency, the study focuses on model-specific evidence availability rather "
        "than clinical sufficiency itself. The paired risk module followed the same analytical "
        "setting: it compared full-view and limited-view predictions to investigate limited-view "
        "underestimation and was not intended to operate on a single image."
    ),
]

CAPTION = (
    "Study overview. (A) Technical gradability and task-relevant evidence availability are distinct: "
    "a clear image may still be evidentially incomplete. (B) Region masks and pixel apertures alter "
    "the composition, area, location, shape, and topology of visible evidence while preserving the "
    "source label. (C) RegionSuff forms a 3×3 evidence bank, extracts frozen retinal image-quality "
    "representations, and normalizes masked attention only across visible regional tokens. (D) The "
    "evidence chain combines internal robustness testing, functional intervention, independent "
    "severity evaluation, and paired-view risk auditing. All view restrictions are synthetic; the "
    "study does not provide clinician-adjudicated sufficiency labels or a deployable single-view "
    "risk head."
)


def copy_run_properties(source_run, target_run):
    if source_run is not None and source_run._element.rPr is not None:
        target_run._element.insert(0, deepcopy(source_run._element.rPr))


def replace_text(paragraph, text):
    source_run = paragraph.runs[0] if paragraph.runs else None
    paragraph.clear()
    run = paragraph.add_run(text)
    copy_run_properties(source_run, run)
    return run


def force_times(run, size=None, bold=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def replace_figure(paragraph):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(FIGURE), width=Inches(5.75))


def replace_caption(paragraph):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = True
    label = paragraph.add_run("Figure 1. ")
    force_times(label, size=10.5, bold=True)
    body = paragraph.add_run(CAPTION)
    force_times(body, size=10.5, bold=False)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)

    document = Document(str(SOURCE))
    paragraphs = document.paragraphs

    # Synchronize the approved three-paragraph abstract and exact-match keywords.
    for index, text in zip((6, 7, 8), ABSTRACT):
        run = replace_text(paragraphs[index], text)
        force_times(run, size=12)

    paragraphs[9].clear()
    key_label = paragraphs[9].add_run("Keywords:")
    force_times(key_label, size=12, bold=True)
    key_body = paragraphs[9].add_run(" " + KEYWORDS)
    force_times(key_body, size=12, bold=False)

    # Synchronize the approved study overview/reporting boundary.
    for index, text in zip((19, 20), STUDY_OVERVIEW):
        run = replace_text(paragraphs[index], text)
        force_times(run, size=12)

    # Replace the old Figure 1 with the enlarged Times New Roman version.
    replace_figure(paragraphs[21])
    replace_caption(paragraphs[22])

    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
