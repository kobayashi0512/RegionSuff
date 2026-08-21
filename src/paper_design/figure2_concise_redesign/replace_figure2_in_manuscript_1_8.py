from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.7.docx')
FIGURE = Path('/Users/tongyue/Documents/ppf/第四论文/图2_架构精简重绘版_核心公式/slide-1.png')
OUTPUT = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.8.docx')

CAPTION = (
    'Figure 2. RegionSuff architecture. (A) A UWF image is divided into a 3×3 regional evidence bank, '
    'and the binary visibility map m identifies the evidence available in a view. (B) Each visible regional '
    'crop is encoded with the shared frozen RetinaRadar EfficientNet-B0 encoder and a trainable projection. '
    'The masked attention pool assigns zero weight to unavailable regions and normalizes weights only across '
    'the remaining regional tokens before producing the three-class DR severity probability. (C) In the '
    'exploratory paired-view audit, the same fitted severity model is evaluated under full and center-plus-cross '
    'masks. The two outputs are compared using probability, severity, and uncertainty differences to rank '
    'limited-view underestimation. This paired analysis requires two predictions and is not a single-image '
    'triage rule.'
)

REPRESENTATION_SENTENCE = (
    'Each retinal box was divided into K=9 rectangular regions. Each crop was resized to 256×256, '
    'center-cropped to 224×224, and normalized with ImageNet channel statistics. A local RetinaRadar '
    'EfficientNet-B0 checkpoint produced a frozen global representation for every region. The checkpoint has '
    '17 quality-related outputs covering laterality, image type, artifacts, clarity, illumination, contrast, '
    'field, and usability; it is not a disease foundation model. Only the projection, attention, and DR '
    'classifier were trained. The shared feature path is summarized in Fig. 2B. The exact local checkpoint '
    'checksum is reported in the code-and-result archive and Reference 23. [18,23]'
)


def replace_embedded_image(document_part, inline_shape, image_path: Path) -> None:
    """Replace Figure 2 image bytes without changing Word layout geometry."""
    blip = inline_shape._inline.graphic.graphicData.pic.blipFill.blip
    relation_id = blip.get(qn('r:embed'))
    if not relation_id:
        raise RuntimeError('The Figure 2 inline image lacks an embedded image relationship.')
    image_part = document_part.related_parts[relation_id]
    image_part._blob = image_path.read_bytes()
    if hasattr(image_part, '_image'):
        image_part._image = None


def set_paragraph_text(paragraph, value: str) -> None:
    paragraph.clear()
    paragraph.add_run(value)


def main() -> None:
    if not SOURCE.exists() or not FIGURE.exists():
        raise FileNotFoundError('The source manuscript or rendered Figure 2 was not found.')

    document = Document(SOURCE)
    if len(document.inline_shapes) < 2:
        raise RuntimeError(f'Expected Figure 2 at inline image index 1; found only {len(document.inline_shapes)} images.')

    # The manuscript is stable: shape 0 is Figure 1 and shape 1 is Figure 2.
    replace_embedded_image(document.part, document.inline_shapes[1], FIGURE)

    caption_found = False
    representation_updated = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith('Figure 2.'):
            set_paragraph_text(paragraph, CAPTION)
            caption_found = True
        elif text.startswith('Each retinal box was divided into K=9 rectangular regions.'):
            set_paragraph_text(paragraph, REPRESENTATION_SENTENCE)
            representation_updated = True

    if not caption_found:
        raise RuntimeError('Could not locate the Figure 2 caption.')
    if not representation_updated:
        raise RuntimeError('Could not locate the regional-representation paragraph.')

    document.save(OUTPUT)
    print({
        'source': str(SOURCE),
        'output': str(OUTPUT),
        'figure_shape_index': 1,
        'caption_updated': caption_found,
        'representation_updated': representation_updated,
        'figure_size_emu': [document.inline_shapes[1].width, document.inline_shapes[1].height],
    })


if __name__ == '__main__':
    main()
