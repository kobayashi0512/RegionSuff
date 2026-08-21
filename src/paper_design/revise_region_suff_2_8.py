from pathlib import Path
import shutil

from docx import Document


SOURCE = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.7.docx')
OUTPUT = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.8.docx')


def insert_before(doc, target, text, style):
    """Insert a styled paragraph immediately before target."""
    new_paragraph = doc.add_paragraph(text, style=style)
    target._p.addprevious(new_paragraph._p)
    return new_paragraph


def main():
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    # Keep the region numbering unambiguous across the manuscript. The paper
    # describes the grid spatially, so the top-left tile is preferable to a
    # software-specific zero-based index here.
    p_mask = next(
        p for p in doc.paragraphs
        if p.text.startswith('During source training, each of the nine regional tokens')
    )
    p_mask.text = (
        'During source training, each of the nine regional tokens was independently masked with probability ρ=0.30. '
        'If an all-empty mask was sampled, the top-left region in row-major order was restored in the locked Exp35/Exp36 implementation; '
        'the event probability was 0.3⁹≈2.0×10⁻⁵ per training image. A sensitivity run restored the geometric center token instead and, '
        'together with validation-list deduplication, produced nearly unchanged test results. Deterministic full-view, center-only, '
        'and center-plus-cross masks were reserved for evaluation.'
    )

    # Separate the dense training/model-selection section from the controls
    # and augmentation experiments so the methodological roles are explicit.
    p_training_heading = next(
        p for p in doc.paragraphs if p.text == '2.10 Training, controls, and model selection'
    )
    p_training_heading.text = '2.10 Training and model selection'

    p_controls = next(
        p for p in doc.paragraphs if p.text.startswith('Single-split controls varied pooling')
    )
    insert_before(doc, p_controls, '2.11 Control experiments and aperture augmentation', 'Heading 3')

    p_controls.text = (
        'Single-split controls were used to assess pooling (mean, max, and attention), token-mask probability '
        '(0, 0.15, 0.30, 0.45, and 0.60), grid size (1×1 through 4×4), and regional representation '
        '(RetinaRadar, ImageNet ResNet-50, 67-dimensional handcrafted statistics, and random ResNet-50). '
        'The 0.30 mask ratio and 3×3 grid were reference configurations inherited from the experimental registry; '
        'the non-monotonic single-split sensitivity does not justify calling them unique optima. A separate five-seed '
        '4×4 run and per-seed paired bootstrap compared 4×4 with 3×3.'
    )

    p_aug = next(
        p for p in doc.paragraphs if p.text.startswith('For the supplementary pixel-aperture training variant')
    )
    p_aug.text = (
        'A supplementary pixel-aperture augmentation variant generated one additional representation per source-training record. '
        'Aperture shape was sampled uniformly from rectangle and circle; fraction was sampled uniformly from 0.45 to 0.80; '
        'and position was sampled uniformly from center, left, right, superior, and inferior. These representations were '
        'concatenated with the original source-training bank, while token masking remained active. Epoch selection again used '
        'full-view source-validation macro-F1, and this variant was evaluated only on four fixed internal stress apertures.'
    )

    # Renumber the downstream subsections after the new controls subsection.
    p_attention_heading = next(p for p in doc.paragraphs if p.text == '2.11 Attention intervention')
    p_attention_heading.text = '2.12 Functional validation of regional attention'
    p_attention = next(
        p for p in doc.paragraphs if p.text.startswith('Attention was interpreted only with an intervention')
    )
    p_attention.text = (
        'Regional attention was interpreted only through an intervention that held the fitted model and reference label fixed. '
        'Each test image was re-evaluated nine times, with one region removed on each pass. For region i, the occlusion loss '
        'was defined as the decrease in full-view macro-F1 after its removal. Spearman correlation then compared the nine mean '
        'full-view attention weights with the nine occlusion losses. This analysis tests functional alignment rather than causal '
        'attribution; the small number of regions limits inferential power. [20]'
    )

    p_external_heading = next(
        p for p in doc.paragraphs if p.text == '2.12 External evaluation and statistical analysis'
    )
    p_external_heading.text = '2.13 External evaluation and statistical analysis'

    doc.save(OUTPUT)
    print(f'Wrote {OUTPUT}')


if __name__ == '__main__':
    main()
