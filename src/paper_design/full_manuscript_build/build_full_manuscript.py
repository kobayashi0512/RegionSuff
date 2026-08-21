#!/usr/bin/env python3
"""Build the complete RegionSufficiency manuscript from locked local results."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/tongyue/Documents/77f /7hao")
BASE_DIR = ROOT / "paper_design/manuscript_build"
sys.path.insert(0, str(BASE_DIR))
import build_manuscript as base  # noqa: E402

REFERENCE = Path("/Users/tongyue/Documents/ppf/第二论文/VOCT_2.0.docx")
FIG = ROOT / "paper_design/figures"
OUTPUT = ROOT / "RegionSufficiency_完整论文_含全套图表公式参考文献.docx"

INK, BROWN, UMBER, ORANGE, OLIVE, ROSE = base.INK, base.BROWN, base.UMBER, base.ORANGE, base.OLIVE, base.ROSE
PAPER, PALE = base.PAPER, base.PALE


def add_text(doc, text, cite=None, first_line=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.65) if first_line else None
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    base.set_run_font(r, size=10.3, color=INK)
    if cite:
        r = p.add_run(cite)
        base.set_run_font(r, size=10.3, color=UMBER)
    base.set_paragraph_keep(p, keep_lines=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet" if "List Bullet" in [s.name for s in doc.styles] else None)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    base.set_run_font(r, size=10.1, color=INK)
    return p


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("When Is a Fundus Image Diagnostically Sufficient?")
    base.set_run_font(r, name="Arial", size=18, bold=True, color=BROWN)
    r.add_break()
    r = p.add_run("Visibility-Conditioned Regional Evidence Aggregation under Limited Fields of View")
    base.set_run_font(r, name="Arial", size=14.5, bold=True, color=UMBER)
    base.add_bottom_rule(p, color=ORANGE, size=12, space=7)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(11)
    r = p.add_run("Anonymous manuscript for peer review")
    base.set_run_font(r, size=9.5, italic=True, color=OLIVE)


def add_table(doc, number, caption, headers, rows, widths=None, font=8.1):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths is None:
        widths = [Cm(16.2 / len(headers))] * len(headers)
    for cell, header, width in zip(table.rows[0].cells, headers, widths):
        cell.width = width
        base.set_cell_shading(cell, BROWN)
        base.set_cell_border(cell, bottom={"val": "single", "sz": "6", "color": ORANGE})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        base.set_run_font(r, name="Arial", size=font, bold=True, color="FFFFFF")
    base.set_repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, (cell, value, width) in enumerate(zip(cells, row, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            base.set_cell_shading(cell, PAPER if ridx % 2 == 0 else PALE)
            base.set_cell_border(cell, bottom={"val": "single", "sz": "3", "color": "E8C6AA"})
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            base.set_run_font(r, size=font, bold=(cidx == 0), color=INK)
        base.set_cant_split(table.rows[-1])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(f"Table {number}. ")
    base.set_run_font(r, size=8.8, bold=True, color=UMBER)
    r = p.add_run(caption)
    base.set_run_font(r, size=8.8, color=INK)
    return table


def add_eq(doc, n, equation, explanation):
    base.add_equation(doc, f"({n})   {equation}", explanation)


def add_reference(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.62)
    p.paragraph_format.first_line_indent = Cm(-0.62)
    p.paragraph_format.space_after = Pt(2.2)
    r = p.add_run(f"{n}. {text}")
    base.set_run_font(r, size=8.7, color=INK)
    base.set_paragraph_keep(p, keep_lines=True)


def build():
    doc = Document(str(REFERENCE))
    base.clear_document_body(doc)
    base.style_document(doc)
    doc.styles["Normal"].font.size = Pt(10.3)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.05
    base.add_header_footer(doc)
    hp = doc.sections[0].header.paragraphs[0]
    hp.text = "REGIONSUFFICIENCY  ·  ANONYMOUS FULL MANUSCRIPT"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        base.set_run_font(run, name="Arial", size=8.5, bold=True, color=UMBER)
    base.add_bottom_rule(hp, color="E8C6AA", size=6, space=2)
    for section in doc.sections:
        section.top_margin = Cm(1.75)
        section.bottom_margin = Cm(1.55)
        section.left_margin = Cm(1.85)
        section.right_margin = Cm(1.85)
    add_title(doc)

    base.add_heading(doc, "Abstract", 1)
    add_text(doc, "Background: Automated fundus-image analysis usually assumes that an acquired image already contains enough evidence for the requested diagnosis. That assumption can fail even when the image is sharp and apparently usable, because informative retinal regions may be absent from the field of view.", first_line=False)
    add_text(doc, "Methods: We formulate diagnostic sufficiency as visibility-conditioned inference. RegionSufficiency divides ultra-widefield (UWF) fundus images into a 3×3 evidence bank, extracts frozen retinal-quality representations, and aggregates only visible tokens using masked attention trained under random region removal. Development used 1,630 UWF-DR images with patient-grouped splits and five random seeds. Independent same-modality replication used the official 2,597-image MMRDR-UWF test subset without target-label training, calibration, early stopping, or model selection. A paired-view risk head ranked limited-view underestimation and was calibrated by a one-sided Clopper–Pearson bound.", first_line=False)
    add_text(doc, "Results: Internal macro-F1 was 0.830±0.017 under the full view, 0.790±0.012 under the center view, and 0.829±0.017 under the center-plus-cross view. External MMRDR-UWF macro-F1 was 0.625±0.010, 0.600±0.004, and 0.642±0.005, respectively. Pixel-space stress tests showed that visible fraction, location, and shape interacted non-monotonically. Occluding the central region caused the largest macro-F1 loss (0.065); regional attention and occlusion loss were directionally aligned (Spearman ρ=0.60, p=0.088). The paired risk head achieved held-out AUROC 0.740; nested grouped calibration yielded AUROC 0.653±0.034 and revealed substantial threshold-transport variability.", first_line=False)
    add_text(doc, "Conclusions: Diagnostic sufficiency is better treated as a property of visible evidence composition than as a synonym for photographic quality or field area. The proposed framework measures this hidden failure mode but does not yet constitute a clinically deployable single-view safety system.", first_line=False)
    base.add_key_value(doc, "Keywords:", "diagnostic sufficiency; ultra-widefield fundus imaging; diabetic retinopathy; limited field of view; masked attention; selective prediction")

    base.add_heading(doc, "1. Introduction", 1, page_break=True)
    add_text(doc, "Diabetic retinopathy (DR) is a major cause of preventable visual loss, and retinal photography has enabled both teleophthalmic screening and autonomous artificial-intelligence systems. Modern classifiers can identify referable DR across large and multiethnic cohorts, but their reported discrimination does not by itself guarantee that an individual acquisition contains all evidence required for the requested judgment. [1–5]")
    add_text(doc, "Most practical pipelines place image-quality assessment before disease inference. They measure focus, illumination, contrast, artifacts, or overall gradability and then accept or reject an image. Large quality datasets and quality-aware diagnosis methods have improved this step. Nevertheless, photographic quality and diagnostic sufficiency are not identical: a focused, well-exposed image can still omit the macula, optic disc, or peripheral lesions needed for a particular task. [6–12]")
    add_text(doc, "This distinction is especially relevant to UWF imaging. Standard color fundus photography usually captures 30–60°; UWF extends toward approximately 200° and can reveal predominantly peripheral DR lesions that are associated with disease progression. A classifier that receives a restricted field can therefore underestimate the full-eye burden even when the visible pixels are technically excellent. [13–16]")
    add_text(doc, "Existing full-image networks, global uncertainty scores, and attention maps do not directly solve this problem. A network generally returns a class for every image. Entropy reports ambiguity, but it does not encode which retinal regions were unavailable. Attention assigns weights to observed features; it does not distinguish low importance from non-observation, and attention should not be equated with causal explanation without intervention. [17–20]")
    add_text(doc, "We introduce RegionSufficiency, a task formulation and regional algorithm for asking whether the currently visible retinal evidence is sufficient for a specified disease judgment. The key operation is visibility-conditioned aggregation: unavailable regions are removed before attention normalization, and random region masking teaches the classifier to operate under evidence loss. Controlled masks vary area, location, shape, and topology while holding the patient and label fixed. An independent UWF dataset then tests whether the observed topology effects survive domain shift.")
    add_text(doc, "Our contributions are fourfold:")
    add_bullet(doc, "We separate diagnostic sufficiency from generic image quality and define a measurable limited-view underestimation endpoint.")
    add_bullet(doc, "We propose a visibility-conditioned regional evidence model with explicit masks, retinal-quality representations, attention aggregation, and stochastic evidence removal.")
    add_bullet(doc, "We audit evidence topology through continuous field size, location, shape, region-bank ablations, and attention occlusion rather than relying on one central crop.")
    add_bullet(doc, "We replicate the main view effects in an independent patient-level UWF test set and report both fixed and nested selective-risk calibration, including negative and unstable findings.")
    base.add_figure(doc, FIG / "Figure1_Study_Overview_RMB20.png", 1, "Study overview. (A) A conventional quality gate can accept a clear image whose disease-relevant evidence is incomplete. (B) UWF-DR supplies development data and MMRDR-UWF supplies independent same-modality replication. (C) RegionSufficiency converts the retina into a visible evidence set and masks unavailable units before aggregation. (D) The study links controlled evidence topology, internal multi-seed testing, external UWF replication, and selective-risk auditing. Geometric views are experimental proxies, not recommendations to crop clinical images.")

    base.add_heading(doc, "2. Materials and Methods", 1, page_break=True)
    base.add_heading(doc, "2.1 Study design, endpoints, and reporting boundary", 2)
    add_text(doc, "This retrospective computational study used de-identified public datasets. No new participant recruitment or image acquisition was performed. The primary endpoint was patient-held-out macro-F1 for three-class DR severity under the center-plus-cross mask. Secondary endpoints were accuracy, ordinal severity mean absolute error (MAE), underestimation rate, and selective accepted underestimation risk. All model-development decisions were made using source validation data. MMRDR-UWF labels were accessed only for final evaluation.")
    add_text(doc, "Claims are deliberately bounded. The study evaluates a model-specific notion of sufficiency for DR severity, not a universal clinical definition. Controlled masks are geometric proxies rather than real paired 45°/UWF photographs. The paired-view risk component uses both full and limited predictions and is therefore an analysis tool, not a deployable single-image triage head.")

    base.add_heading(doc, "2.2 Datasets and label harmonization", 2)
    add_table(doc, 1, "Datasets, label mapping, and strictly separated experimental roles.",
              ["Dataset", "Images", "Patients/split", "Labels used here", "Role and leakage control"],
              [
                  ["UWF-DR", "1,630", "809 patients; train 876, validation 364, test 390", "Normal 496; NPDR 634; PDR 500", "Development, ablation, controlled masks; patient grouping"],
                  ["MMRDR-UWF", "2,597", "Official patient-level test subset", "Grade 0→Normal (800); grades 1–3→NPDR (1,411); grade 4→PDR (386)", "Independent UWF replication; no target-label training or selection"],
              ], [Cm(2.5), Cm(1.4), Cm(3.7), Cm(4.0), Cm(4.6)], font=7.7)
    add_text(doc, "UWF-DR contains 1,630 scanning-laser UWF images released with patient identity and Normal/NPDR/PDR labels. The patient-grouped development split contained 876 training images (260/351/265), 364 validation images (107/146/111), and 390 test images (129/137/124). [21]")
    add_text(doc, "MMRDR contains CFP, OCT, and UWF modalities. We used only the official UWF testing split because it was partitioned at patient level by the dataset authors. Its original five-level DR grades were mapped to the source three-class task: grade 0 to Normal, grades 1–3 to NPDR, and grade 4 to PDR. The external set contained 800/1,411/386 images after mapping. [22]")

    base.add_heading(doc, "2.3 Controlled field-of-view generation", 2)
    add_text(doc, "Two complementary perturbation families were constructed from each UWF image. Region-level perturbations selected subsets of a 3×3 evidence bank: full (nine regions), center (one region), and center-plus-cross (center plus its four axial neighbors). Pixel-level perturbations retained a centered rectangle or circle with linear fractions 0.35, 0.45, 0.55, 0.65, 0.75, or 0.90; translated a 0.55 rectangle left, right, superiorly, or inferiorly; and compared shapes at a matched linear fraction. Blacked-out pixels were applied before feature extraction.")
    add_text(doc, "These perturbations do not reproduce optical differences between cameras. Their purpose is causal stress testing: the underlying patient and ground-truth severity remain fixed while the observable evidence changes.")

    base.add_heading(doc, "2.4 Region bank and frozen retinal representation", 2)
    add_text(doc, "After retinal foreground normalization, each image was partitioned into K=9 rectangular regions with row-major indices. Let rᵢ denote the i-th region and bᵢ=(xᵢ¹,yᵢ¹,xᵢ²,yᵢ²) its normalized box. A released RetinaRadar EfficientNet-B0 quality-assessment checkpoint supplied a fixed retinal representation. It predicts laterality, image type, artifacts, clarity, illumination, contrast, field, and usability; it is not a disease model. Only the regional projection, attention aggregator, and DR classifier were trained. EfficientNet-B0 was selected because the local checkpoint was available and frozen; the encoder ablation compares this representation with ImageNet, handcrafted, and random controls. [18,23,24]")
    add_eq(doc, 1, "R(x) = {rᵢ = Crop(x,bᵢ)}ᵢ₌₁ᴷ,   K = 9", "R(x) is the region bank. Crop extracts a fixed spatial unit using normalized box bᵢ; the patient image x and label y do not change across masks.")
    add_eq(doc, 2, "zᵢ = P(E(rᵢ)) ∈ ℝᵈ", "E is the frozen RetinaRadar EfficientNet-B0 encoder; P is the trainable projection into a shared d-dimensional token space.")

    base.add_heading(doc, "2.5 Visibility-conditioned attention", 2)
    add_text(doc, "For a view mask m∈{0,1}ᴷ, mᵢ=1 means region i is available. The score network first evaluates every regional token, after which masking occurs before softmax normalization. Therefore an unavailable region receives exactly zero weight and cannot affect the representation.")
    add_eq(doc, 3, "eᵢ = wₐᵀ tanh(Wₐzᵢ + bₐ)", "eᵢ is the unnormalized evidence score for region i; Wₐ, bₐ, and wₐ are learned attention parameters.")
    add_eq(doc, 4, "aᵢ(m) = mᵢ exp(eᵢ) / [Σⱼ₌₁ᴷ mⱼ exp(eⱼ) + ε]", "aᵢ(m) is masked attention. Multiplication by mᵢ removes invisible regions before normalization; ε prevents division by zero.")
    add_eq(doc, 5, "h(m) = Σᵢ₌₁ᴷ aᵢ(m)zᵢ", "h(m) is the image-level representation formed only from visible regional evidence.")
    add_eq(doc, 6, "p(m) = softmax(Wch(m)+bc)", "p(m)=[p₀,p₁,p₂] contains class probabilities for Normal, NPDR, and PDR.")
    add_eq(doc, 7, "Lcls = −(1/N)ΣₙΣc wc 1[yₙ=c] log pₙ,c(mₙ)", "Class-weighted cross-entropy trains the projection, attention, and classifier while the encoder remains frozen.")

    base.add_heading(doc, "2.6 Stochastic evidence masking", 2)
    add_text(doc, "During training, visible regions were independently removed with probability ρ=0.30, subject to retaining at least one region. This augmentation changes evidence availability rather than disease labels and prevents the aggregator from depending on one permanently visible token.")
    add_eq(doc, 8, "m̃ᵢ ~ Bernoulli(1−ρ),   ρ=0.30,   conditioned on Σᵢm̃ᵢ≥1", "m̃ is the stochastic training mask. The final 30% rate was prespecified in the locked experiment registry; sensitivity from 0 to 60% is reported.")
    base.add_figure(doc, FIG / "Figure2_Method_Architecture_RMB20.png", 2, "RegionSufficiency architecture. (A) UWF images are divided into a 3×3 evidence bank; deterministic evaluation masks and 30% random training masks define visibility. (B) A shared frozen RetinaRadar EfficientNet-B0 encoder and trainable projection produce regional tokens. (C) Masked attention aggregates only visible tokens for Normal/NPDR/PDR prediction. (D) Paired full/limited predictions form an 18-dimensional evidence-gap vector for an underestimation-risk model. The risk component is an audit that requires paired views.")

    base.add_heading(doc, "2.7 Disease severity and underestimation", 2)
    add_eq(doc, 9, "ŷ(m)=argmaxc pc(m);   s(m)=Σc₌₀² c·pc(m)", "ŷ is the categorical prediction and s(m) is expected ordinal severity on the 0–2 scale.")
    add_eq(doc, 10, "u(mL,y)=1[ŷ(mL)<y]", "The safety event u equals one when the limited-view prediction is less severe than the reference label.")
    add_eq(doc, 11, "MAE=(1/N)Σₙ|ŷₙ−yₙ|;   Under=(1/N)Σₙuₙ", "MAE penalizes ordinal distance; Under is the proportion of cases whose disease severity is underestimated.")
    add_eq(doc, 12, "Macro-F1=(1/C)Σc 2·Precisionc·Recallc/(Precisionc+Recallc)", "Macro-F1 weights each class equally and is the primary endpoint because the three severity classes are not equally frequent.")

    base.add_heading(doc, "2.8 Paired-view evidence-gap risk", 2)
    add_text(doc, "The paired risk audit compares probability vectors under full (pF) and limited (pL) evidence. In addition to their direct difference, the feature vector contains expected severity, entropy, and symmetric Kullback–Leibler divergence. This design asks whether the model's own evidence shift predicts a clinically undesirable downward severity error.")
    add_eq(doc, 13, "H(p)=−Σc pc log(pc+ε)", "H(p) is predictive entropy; larger values indicate a more diffuse class distribution.")
    add_eq(doc, 14, "DSKL(pF,pL)=½[KL(pF∥pL)+KL(pL∥pF)]", "Symmetric KL quantifies the magnitude of the full-to-limited distribution shift without privileging either direction.")
    add_eq(doc, 15, "φ=[pF,pL,pF−pL,|pF−pL|,sF,sL,sF−sL,HF,HL,DSKL]∈ℝ¹⁸", "The 18 dimensions are 12 probability features, three severity features, two entropies, and one symmetric-KL feature.")
    add_eq(doc, 16, "q(φ)=σ(βᵀφ+b)", "A logistic model fitted on validation evidence-gap features estimates the probability q of limited-view underestimation.")

    base.add_heading(doc, "2.9 Risk-controlled selective acceptance", 2)
    add_text(doc, "At target risk α, samples were sorted by predicted risk q. The largest threshold was selected whose one-sided 95% exact binomial upper bound did not exceed α on calibration data. Samples with q≤τα were accepted and the remainder deferred. This is a risk–coverage trade-off rather than probability calibration or a generic misclassification/OOD score alone. [25–27,29,30]")
    add_eq(doc, 17, "τₐ=max{τ : UCP(kτ,nτ;0.95)≤α};   accept ⇔ q≤τₐ", "nτ and kτ are accepted calibration cases and accepted underestimations; UCP is the one-sided Clopper–Pearson upper confidence bound.")
    add_eq(doc, 18, "Coverage(τ)=nτ/N;   Risk(τ)=kτ/nτ", "Coverage measures retained workload; selective risk measures underestimation among accepted cases.")

    base.add_heading(doc, "2.10 Training, model selection, and ablations", 2)
    add_text(doc, "Feature normalization was fitted on source training data only. Validation macro-F1 under the center-plus-cross mask controlled early stopping and model selection. The held-out internal test set was then evaluated across seeds 0–4. The frozen registry selected 3×3 regions, RetinaRadar features, attention pooling, and ρ=0.30. If validation scores tied within 0.005, the smaller model and lower validation underestimation were preferred.")
    add_text(doc, "Ablations changed one component at a time: grid granularity (1×1, 2×2, 3×3, 4×4), pooling (mean, max, attention), training mask ratio (0, 0.15, 0.30, 0.45, 0.60), and regional representation (RetinaRadar, ImageNet ResNet-50, 67-dimensional handcrafted statistics, random ResNet-50). A separate five-seed 4×4 experiment and paired bootstrap assessed whether the finer grid consistently improved the prespecified 3×3 model.")

    base.add_heading(doc, "2.11 Attention intervention", 2)
    add_text(doc, "Attention was not accepted as an explanation solely because a heatmap appeared plausible. Each of the nine regions was removed in turn from every held-out image while model weights and labels remained fixed. Region importance was defined as the macro-F1 loss after occlusion, and Spearman correlation compared the rankings of mean attention and intervention loss. [19,20]")
    add_eq(doc, 19, "Δi = Macro-F1full − Macro-F1mask(i)", "Δi is the causal performance loss caused by hiding region i. Positive values indicate that the region contributed useful held-out evidence.")

    base.add_heading(doc, "2.12 External evaluation and statistical analysis", 2)
    add_text(doc, "All five source-trained 3×3 models were applied to the MMRDR-UWF official test split without MMRDR training, calibration, early stopping, or model selection. Point metrics were summarized as mean±standard deviation across source seeds. View comparisons used 2,000 paired image-level bootstrap resamples, preserving the same external images for each view within a seed. [28]")
    add_eq(doc, 20, "Δ*b = T({(yjb,ŷL,jb)}) − T({(yjb,ŷF,jb)}),   b=1,…,B", "For each bootstrap b, the same sampled image indices jb are used for limited and full views. Percentiles of Δ*b give paired 95% intervals.")
    add_text(doc, "The exploratory MMRDR quality proxy combined retinal coverage, within-retina contrast, and first-difference sharpness. It had no clinician quality label and is reported only as a domain-shift diagnostic. A four-fold outer GroupKFold analysis on source train+validation patients further evaluated selective-threshold transport: each outer-training fold was divided into independent model-fit and calibration groups, and the original source test set remained untouched.")

    base.add_heading(doc, "3. Results", 1, page_break=True)
    base.add_heading(doc, "3.1 Patient-grouped internal performance", 2)
    add_text(doc, "Across five random seeds, full-view macro-F1 was 0.830±0.017. Restricting evidence to the center reduced macro-F1 to 0.790±0.012 and increased underestimation from 0.070±0.007 to 0.093±0.006. Center-plus-cross recovered macro-F1 to 0.829±0.017, severity MAE to 0.182±0.016, and underestimation to 0.075±0.012 (Table 2). The center-plus-cross topology therefore retained most full-view discrimination despite exposing only five of nine region tokens.")
    add_table(doc, 2, "Internal held-out results across five patient-grouped random seeds. Values are mean±SD.",
              ["View", "Accuracy", "Macro-F1", "Severity MAE", "Underestimation"],
              [
                  ["Full", "0.829±0.017", "0.830±0.017", "0.180±0.018", "0.070±0.007"],
                  ["Center", "0.789±0.012", "0.790±0.012", "0.228±0.013", "0.093±0.006"],
                  ["Center + cross", "0.828±0.017", "0.829±0.017", "0.182±0.016", "0.075±0.012"],
              ], [Cm(3.7), Cm(3.0), Cm(3.0), Cm(3.0), Cm(3.5)], font=8.3)

    base.add_heading(doc, "3.2 Area, location, shape, and topology are not interchangeable", 2)
    add_text(doc, "Pixel-space experiments produced an overall rise in macro-F1 as visible fraction approached 0.90, but the trajectory was non-monotonic and highly variable at small fractions (Figure 3A). For rectangular masks, macro-F1 changed from 0.643±0.150 at fraction 0.35 to 0.814±0.022 at 0.90; for circles it changed from 0.584±0.114 to 0.819±0.018. At fraction 0.55, the circle reached 0.750±0.046 whereas the rectangle reached 0.678±0.085, showing that equal linear extent did not imply equal evidence topology or visible pixel area.")
    add_text(doc, "Translating the 0.55 rectangle yielded similar mean macro-F1 values (0.673–0.675 outside the center) but large seed variability and different underestimation behavior: inferior translation produced underestimation 0.159±0.099 compared with 0.050±0.017 for the centered window. Thus location effects were more visible in safety-oriented errors than in the mean F1 alone. Region-bank masks showed the clearest topology result: center-plus-cross consistently recovered the loss caused by the center-only mask.")
    base.add_figure(doc, FIG / "Figure3_Controlled_Visibility_RMB20.png", 3, "Controlled visibility tests. (A) Five-seed macro-F1 across continuous centered rectangle and circle fractions. (B) Macro-F1 for fixed-size translated rectangles. (C) Matched linear fraction with different shape; shape changes topology and visible pixel area. (D) Paired five-seed region-bank results. Error bars denote SD. These experiments are mechanism stress tests and not clinical camera equivalence studies.")

    base.add_heading(doc, "3.3 Ablations support retinal regional representation and attention, but not a single monotonic hyperparameter optimum", 2)
    add_text(doc, "With the same 3×3 attention head, the frozen RetinaRadar representation achieved center-plus-cross macro-F1 0.837, compared with 0.772 for ImageNet ResNet-50, 0.604 for handcrafted regional statistics, and 0.161 for random ResNet-50. This control indicates that retinal pretraining, rather than the small attention head alone, supplied useful regional features (Figure 4A).")
    add_text(doc, "Attention pooling improved full-view macro-F1 from 0.799 (mean) and 0.810 (max) to 0.840, while center-plus-cross performance was 0.820, 0.824, and 0.837, respectively. Mask-ratio sensitivity was non-monotonic; ρ=0.30 was retained because it was locked before final multi-seed reporting, not because it was the largest single-split test value. The 4×4 model averaged 0.842±0.016 under center-plus-cross, compared with 0.829±0.017 for 3×3, but every per-seed paired 95% macro-F1 interval crossed zero. The 3×3 model was therefore retained as the prespecified and more interpretable configuration.")
    add_table(doc, 3, "Single-split component ablations; test values are descriptive and did not override the locked final configuration.",
              ["Component", "Variant", "Full macro-F1", "Center macro-F1", "Center + cross macro-F1"],
              [
                  ["Pooling", "Mean", "0.799", "0.759", "0.820"], ["", "Max", "0.810", "0.778", "0.824"], ["", "Attention", "0.840", "0.793", "0.837"],
                  ["Mask ratio", "0.00 / 0.15 / 0.30 / 0.45 / 0.60", "0.837 / 0.817 / 0.840 / 0.834 / 0.853", "0.801 / 0.795 / 0.793 / 0.795 / 0.798", "0.847 / 0.820 / 0.837 / 0.843 / 0.856"],
                  ["Encoder", "RetinaRadar / ImageNet / handcrafted / random", "0.840 / 0.780 / 0.605 / 0.161", "—", "0.837 / 0.772 / 0.604 / 0.161"],
              ], [Cm(2.4), Cm(5.1), Cm(2.9), Cm(2.9), Cm(3.4)], font=7.4)
    base.add_figure(doc, FIG / "Figure4_Ablations_RMB20.png", 4, "Ablations and controls. (A) Frozen regional representation comparison. (B) Mean, max, and attention aggregation. (C) Training-mask sensitivity with the prespecified 0.30 configuration marked. (D) Per-seed paired bootstrap differences between 4×4 and 3×3 under center-plus-cross; all 95% intervals cross zero.")

    base.add_heading(doc, "3.4 Attention aligns directionally with interventional importance", 2)
    add_text(doc, "The middle-center region received mean attention 0.398 and caused the largest macro-F1 loss when occluded (0.065). Removing middle-left caused loss 0.018 and middle-right loss 0.011. Several peripheral regions produced near-zero or slightly negative loss, showing that learned attention was not simply a smooth centrality prior. Across nine regions, attention rank and occlusion-loss rank had Spearman ρ=0.60 (p=0.088). The direction is mechanistically supportive, but the region count is too small for a definitive statistical claim.")
    add_text(doc, "Class-stratified held-out attention remained center-dominant for Normal and NPDR. PDR shifted more weight toward the middle-left and middle-right regions while retaining a central maximum. This pattern is compatible with distributed severe-disease evidence, but it is an aggregation analysis rather than lesion localization because no lesion-level labels entered training.")
    base.add_figure(doc, FIG / "Figure5_Attention_Occlusion_RMB20.png", 5, "Attention intervention. (A) Mean regional attention. (B) Macro-F1 loss after independently hiding each region. (C) Attention versus intervention loss, with the observed Spearman association. (D) Mean attention stratified by true held-out class. Attention is interpreted only together with the occlusion intervention.")

    base.add_heading(doc, "3.5 Independent MMRDR-UWF replication", 2)
    add_text(doc, "The source-trained models showed an expected domain-shift decrease on MMRDR-UWF, but the relative evidence topology replicated. Full-view external macro-F1 was 0.625±0.010; center fell to 0.600±0.004; center-plus-cross increased to 0.642±0.005. Center-plus-cross also reduced severity MAE from 0.400±0.013 to 0.378±0.006 and underestimation from 0.272±0.013 to 0.260±0.005 (Table 4).")
    add_text(doc, "Paired bootstrap intervals for center-versus-full macro-F1 were predominantly negative or touched zero. Center-plus-cross was positive in all five seeds, with seed-specific point differences from 0.007 to 0.025. This counter-intuitive advantage over the nominal full view is not interpreted as evidence that cropping improves clinical diagnosis. It may reflect suppression of peripheral device artifacts, different framing, or a favorable evidence-to-noise ratio under domain shift.")
    add_text(doc, "The unlabeled external quality proxy did not induce a uniform monotonic performance gradient across seeds. The highest proxy quartile improved over the lowest in four seeds but changed little in one. This reinforces the need for clinician-adjudicated quality and sufficiency labels rather than relying on handcrafted quality proxies.")
    add_table(doc, 4, "Independent MMRDR-UWF results across five source-trained seeds. Values are mean±SD; no MMRDR labels entered training or model selection.",
              ["View", "Accuracy", "Macro-F1", "Severity MAE", "Underestimation", "Δ macro-F1 vs full"],
              [
                  ["Full", "0.627±0.010", "0.625±0.010", "0.400±0.013", "0.272±0.013", "Reference"],
                  ["Center", "0.602±0.004", "0.600±0.004", "0.432±0.005", "0.284±0.016", "−0.025±0.013"],
                  ["Center + cross", "0.644±0.004", "0.642±0.005", "0.378±0.006", "0.260±0.005", "+0.017±0.006"],
              ], [Cm(2.8), Cm(2.4), Cm(2.5), Cm(2.5), Cm(2.7), Cm(3.0)], font=7.9)
    base.add_figure(doc, FIG / "Figure6_External_Replication_RMB20.png", 6, "Independent same-modality UWF replication. (A) Internal and external macro-F1. (B) External ordinal error and underestimation. (C) Per-seed paired 2,000-resample bootstrap differences versus full view. (D) Exploratory lowest/highest quartiles of an unlabeled external quality proxy. External labels were used only for evaluation.")

    base.add_heading(doc, "3.6 Paired-view selective underestimation risk", 2)
    add_text(doc, "On the fixed held-out internal test set, the 18-dimensional risk head achieved AUROC 0.740 and average precision 0.139 at an event prevalence of 0.077. Severity gap alone reached AUROC 0.677, whereas symmetric KL alone reached 0.539. At a target accepted risk of 5%, the calibrated threshold retained 62.8% of cases and observed accepted underestimation 2.9%; at targets of 10% and 20%, coverage was 91.8% and 99.7% with observed risk 7.5% and 7.7%.")
    add_text(doc, "The nested grouped audit was less optimistic. Mean risk AUROC was 0.653±0.034. At target 5%, accepted coverage was 56.7%±27.4% and accepted risk 5.0%±3.3%; individual folds ranged from 36.5% to 95.8% coverage and 0.8% to 8.7% accepted risk. At target 10%, accepted risk was 9.2%±3.4%. These results show that threshold transport, not ranking alone, is the main barrier to a stable deployment policy.")
    add_table(doc, 5, "Selective underestimation-risk results. Fixed values use the locked validation/test split; nested values are mean±SD across four grouped outer folds.",
              ["Protocol", "Risk AUROC", "Target", "Accepted coverage", "Accepted underestimation risk"],
              [
                  ["Fixed held-out", "0.740", "5%", "0.628", "0.029"], ["", "", "10%", "0.918", "0.075"], ["", "", "20%", "0.997", "0.077"],
                  ["Nested grouped", "0.653±0.034", "5%", "0.567±0.274", "0.050±0.033"], ["", "", "10%", "0.812±0.165", "0.092±0.034"], ["", "", "20%", "0.980±0.014", "0.097±0.040"],
              ], [Cm(3.6), Cm(3.0), Cm(2.2), Cm(3.5), Cm(4.0)], font=8.0)
    base.add_figure(doc, FIG / "Figure7_Selective_Risk_RMB20.png", 7, "Selective risk. (A) Underestimation-risk discrimination for the paired model and direct scores. (B) Coverage retained at increasing risk targets. (C) Observed accepted risk versus target. (D) Nested outer-fold coverage–risk points. Error bars denote SD across outer folds. The paired risk head is an evidence-loss audit, not a deployable single-view model.")

    base.add_heading(doc, "4. Discussion", 1, page_break=True)
    base.add_heading(doc, "4.1 Principal findings", 2)
    add_text(doc, "This study addresses a hidden assumption in retinal AI: that an image deemed technically acceptable also contains sufficient evidence for the requested diagnosis. RegionSufficiency makes evidence availability explicit. Its central empirical result is not that a particular crop is universally optimal; it is that field area, location, shape, and topology produce different disease-model behavior even when the patient and label are unchanged.")
    add_text(doc, "Three observations support this conclusion. First, center-only evidence consistently degraded internal and external performance, whereas center-plus-cross recovered or exceeded the nominal full-view result. Second, retinal-quality regional features materially outperformed ImageNet, handcrafted, and random controls. Third, learned attention received directional support from an intervention: the region with the highest mean attention caused the largest held-out performance loss when removed.")

    base.add_heading(doc, "4.2 Relationship to image quality and quality-aware diagnosis", 2)
    add_text(doc, "Image-quality assessment remains necessary: blur, illumination, contrast, and artifacts can invalidate disease predictions. Our proposal is complementary rather than competing. A quality model asks whether visible structures are captured clearly; a sufficiency model asks whether the visible structures include enough task-relevant evidence. Recent quality-aware systems condition diagnosis on degradation, but the visibility mask provides a separate semantic variable—what was observed—that cannot be reconstructed from sharpness or entropy alone. [6–12]")
    add_text(doc, "This distinction suggests a two-axis acquisition policy. An image can be high-quality and sufficient, high-quality but insufficient, low-quality yet partially informative, or both low-quality and insufficient. Future prospective labels should ask graders to annotate these axes separately and specify the target decision, because sufficiency for referable-DR screening may differ from sufficiency for fine five-level grading or lesion surveillance.")

    base.add_heading(doc, "4.3 Why the external center-plus-cross view exceeded the full view", 2)
    add_text(doc, "The external result requires caution. Under source-to-target shift, the nominal full image contains both disease evidence and device-specific nuisance variation. A structured mask can remove peripheral artifacts or border geometry and increase the signal-to-noise ratio seen by a source-trained model. Thus center-plus-cross outperforming full view is plausible as a model-domain interaction, but it is not evidence that clinicians should discard peripheral information. The correct next experiment is a prospective same-eye, multi-device study with clinician sufficiency labels and explicit foreground/border controls.")

    base.add_heading(doc, "4.4 Interpretability and selective prediction", 2)
    add_text(doc, "Attention visualization is often overinterpreted. By pairing the heatmap with one-region occlusion, we test whether high-weight regions are also functionally important. The observed ρ=0.60 is encouraging but not conclusive, and the p value is reported without dichotomization. Lesion-level masks would permit stronger tests such as whether attention and intervention localize neovascularization or peripheral hemorrhage.")
    add_text(doc, "Selective prediction provides a natural response to insufficiency: accept low-risk cases and defer the remainder. The fixed split showed favorable risk–coverage behavior, but nested grouped calibration exposed large variation. This gap is important. A useful risk ranker is not automatically a safe operational threshold. Prospective deployment would require single-view features, independent calibration, subgroup auditing, temporal monitoring, and a predefined recapture/referral action. [25–27]")

    base.add_heading(doc, "4.5 Strengths", 2)
    add_text(doc, "The study uses patient grouping, a locked primary endpoint, five random seeds, controlled within-image perturbations, encoder/pooling/grid/mask controls, intervention-based attention validation, an independent patient-level UWF test set, paired bootstrap inference, and both fixed and nested selective-risk analyses. Equally important, unfavorable findings—non-monotonic mask sensitivity, unstable continuous-mask seeds, and nested threshold variability—are reported rather than hidden.")

    base.add_heading(doc, "5. Limitations", 1)
    add_text(doc, "First, no dataset provides a clinician-adjudicated label for diagnostic sufficiency. Ground truth is disease severity, and sufficiency is inferred from model behavior under evidence loss. Second, geometric masks do not reproduce optics, centering, eyelash artifacts, or paired-device acquisition. Third, the main task collapses five DR levels into three classes to harmonize datasets. Fourth, RetinaRadar is a locally released quality encoder with incomplete public provenance; it is used frozen and must not be described as a disease foundation model. Fifth, only one independent same-modality UWF test set is used. Sixth, the quality-proxy strata are not clinical quality labels. Seventh, paired-view risk requires the full view and therefore cannot triage a lone limited acquisition. Eighth, bootstrap intervals quantify uncertainty within the available test set and do not replace prospective multicenter validation.")

    base.add_heading(doc, "6. Conclusion", 1)
    add_text(doc, "A fundus image can be visually clear yet diagnostically insufficient. RegionSufficiency converts this overlooked failure mode into a visibility-conditioned regional inference problem. Internal and independent UWF experiments show that evidence topology matters beyond visible area alone, and intervention supports the central learned evidence pattern. Paired selective-risk results identify a path toward deferral, while nested calibration shows why prospective, single-view validation remains necessary.")

    base.add_heading(doc, "Declarations", 1)
    base.add_heading(doc, "Ethics statement", 2)
    add_text(doc, "This work is a secondary computational analysis of de-identified public datasets and involved no new participant recruitment, intervention, or data collection. Ethical approvals and consent/waiver procedures for the source cohorts are reported in the original dataset publications. [21,22]", first_line=False)
    base.add_heading(doc, "Data availability", 2)
    add_text(doc, "UWF-DR is available from Figshare (https://doi.org/10.6084/m9.figshare.31259494). MMRDR is available from Figshare (https://doi.org/10.6084/m9.figshare.29423747.v2). Reuse must follow the source licenses and data terms.", first_line=False)
    base.add_heading(doc, "Code and result availability", 2)
    add_text(doc, "The project archive accompanying this manuscript contains the analysis scripts, frozen experiment registry, machine-readable results, paired-bootstrap outputs, figure-generation source, and editable vector figures. No original experiment output was overwritten during manuscript assembly.", first_line=False)
    base.add_heading(doc, "Competing interests", 2)
    add_text(doc, "The authors declare no competing interests.", first_line=False)

    base.add_heading(doc, "References", 1, page_break=True)
    refs = [
        "Wilkinson CP, Ferris FL III, Klein RE, et al. Proposed international clinical diabetic retinopathy and diabetic macular edema disease severity scales. Ophthalmology. 2003;110:1677–1682. doi:10.1016/S0161-6420(03)00475-5.",
        "Wong TY, Cheung CMG, Larsen M, Sharma S, Simó R. Diabetic retinopathy. Nat Rev Dis Primers. 2016;2:16012. doi:10.1038/nrdp.2016.12.",
        "Gulshan V, Peng L, Coram M, et al. Development and validation of a deep learning algorithm for detection of diabetic retinopathy in retinal fundus photographs. JAMA. 2016;316:2402–2410. doi:10.1001/jama.2016.17216.",
        "Ting DSW, Cheung CYL, Lim G, et al. Development and validation of a deep learning system for diabetic retinopathy and related eye diseases using retinal images from multiethnic populations with diabetes. JAMA. 2017;318:2211–2223. doi:10.1001/jama.2017.18152.",
        "Abràmoff MD, Lavin PT, Birch M, Shah N, Folk JC. Pivotal trial of an autonomous AI-based diagnostic system for detection of diabetic retinopathy in primary care offices. npj Digit Med. 2018;1:39. doi:10.1038/s41746-018-0040-6.",
        "Saha SK, Fernando B, Cuadros J, Xiao D, Kanagasingam Y. Automated quality assessment of colour fundus images for diabetic retinopathy screening in telemedicine. J Digit Imaging. 2018;31:869–878. doi:10.1007/s10278-018-0084-9.",
        "Fu H, Wang B, Shen J, et al. Evaluation of retinal image quality assessment networks in different color-spaces. In: MICCAI 2019. LNCS 11764:48–56. doi:10.1007/978-3-030-32239-7_6.",
        "Gonçalves MB, Nakayama LF, Ferraz D, et al. Image quality assessment of retinal fundus photographs for diabetic retinopathy in the machine learning era: a review. Eye. 2024;38:426–433. doi:10.1038/s41433-023-02717-3.",
        "Jin K, Gao Z, Jiang X, et al. MSHF: A multi-source heterogeneous fundus dataset for image quality assessment. Sci Data. 2023;10:286. doi:10.1038/s41597-023-02188-x.",
        "Liu R, Wang X, Wu Q, et al. DeepDRiD: Diabetic retinopathy-grading and image quality estimation challenge. Patterns. 2022;3:100512. doi:10.1016/j.patter.2022.100512.",
        "Che H, Chen S, Chen H. Image quality-aware diagnosis via meta-knowledge co-embedding. Proc IEEE/CVF Conf Comput Vis Pattern Recognit. 2023.",
        "Oh R, Park UC, Park KH, Park SJ, Yoon CK. Deep learning-based automatic image quality assessment in ultra-widefield fundus photographs. BMJ Open. 2025;15:e100058. doi:10.1136/bmjopen-2025-100058.",
        "Silva PS, Cavallerano JD, Haddad NMN, et al. Peripheral lesions identified on ultrawide field imaging predict increased risk of diabetic retinopathy progression over 4 years. Ophthalmology. 2015;122:949–956. doi:10.1016/j.ophtha.2015.01.008.",
        "Ashrafkhorasani M, Habibi A, Nittala MG, et al. Peripheral retinal lesions in diabetic retinopathy on ultra-widefield imaging. Saudi J Ophthalmol. 2024;38:123–131. doi:10.4103/sjopt.sjopt_151_23.",
        "Rajalakshmi R, Mohammed R, Vengatesan K, et al. Wide-field imaging with smartphone based fundus camera: grading severity and locating peripheral lesions in diabetic retinopathy. Eye. 2024. doi:10.1038/s41433-024-02928-2.",
        "Rajalakshmi R, Prathiba V, Arulmalar S, Usha M. Review of retinal cameras for global coverage of diabetic retinopathy screening. Eye. 2021;35:162–172.",
        "He K, Zhang X, Ren S, Sun J. Deep residual learning for image recognition. Proc IEEE Conf Comput Vis Pattern Recognit. 2016:770–778. doi:10.1109/CVPR.2016.90.",
        "Tan M, Le QV. EfficientNet: Rethinking model scaling for convolutional neural networks. Proc 36th Int Conf Mach Learn. 2019;97:6105–6114.",
        "Ilse M, Tomczak JM, Welling M. Attention-based deep multiple instance learning. Proc 35th Int Conf Mach Learn. 2018;80:2127–2136.",
        "Jain S, Wallace BC. Attention is not Explanation. Proc NAACL-HLT. 2019:3543–3556. doi:10.18653/v1/N19-1357.",
        "Peng S, Yang S, Zhao X, et al. A fundus image dataset for intelligent diabetic retinopathy system. Sci Data. 2026. doi:10.1038/s41597-026-07093-7. Dataset doi:10.6084/m9.figshare.31259494.",
        "Tang Z, Wang L, Guo Z, et al. A multimodal retinal image dataset for diabetic retinopathy detection using foundation models. Sci Data. 2026;13:639. doi:10.1038/s41597-026-07005-9.",
        "RetinaRadar. Multi-label retinal image quality assessment, EfficientNet-B0 checkpoint and local model card. Software/model release, 2025. Apache-2.0.",
        "Deng J, Dong W, Socher R, Li LJ, Li K, Fei-Fei L. ImageNet: A large-scale hierarchical image database. Proc IEEE Conf Comput Vis Pattern Recognit. 2009:248–255. doi:10.1109/CVPR.2009.5206848.",
        "Geifman Y, El-Yaniv R. Selective classification for deep neural networks. Adv Neural Inf Process Syst. 2017;30:4878–4887.",
        "Geifman Y, El-Yaniv R. SelectiveNet: A deep neural network with an integrated reject option. Proc 36th Int Conf Mach Learn. 2019;97:2151–2159.",
        "Clopper CJ, Pearson ES. The use of confidence or fiducial limits illustrated in the case of the binomial. Biometrika. 1934;26:404–413. doi:10.1093/biomet/26.4.404.",
        "Efron B, Tibshirani RJ. An Introduction to the Bootstrap. New York: Chapman & Hall/CRC; 1993.",
        "Guo C, Pleiss G, Sun Y, Weinberger KQ. On calibration of modern neural networks. Proc 34th Int Conf Mach Learn. 2017;70:1321–1330.",
        "Hendrycks D, Gimpel K. A baseline for detecting misclassified and out-of-distribution examples in neural networks. Proc Int Conf Learn Represent. 2017.",
    ]
    for i, ref in enumerate(refs, 1): add_reference(doc, i, ref)

    doc.core_properties.title = "When Is a Fundus Image Diagnostically Sufficient?"
    doc.core_properties.subject = "Visibility-conditioned regional evidence aggregation for limited-view fundus imaging"
    doc.core_properties.keywords = "fundus imaging, diagnostic sufficiency, UWF, regional attention, selective prediction"
    doc.core_properties.comments = "Complete evidence-traceable manuscript assembled from locked local experiment outputs."
    doc.core_properties.author = "Anonymous"
    doc.core_properties.last_modified_by = "Anonymous"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
