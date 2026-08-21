from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement


SOURCE = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.5.docx')
OUTPUT = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.6.docx')


def clear_paragraph(paragraph):
    p = paragraph._p
    p_pr = p.pPr
    for child in list(p):
        if child is not p_pr:
            p.remove(child)


def set_paragraph_text(paragraph, text):
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def m_el(tag):
    return OxmlElement(f'm:{tag}')


def math_run(text):
    run = m_el('r')
    text_node = m_el('t')
    text_node.text = text
    run.append(text_node)
    return run


def math_sub(base, sub):
    node = m_el('sSub')
    e = m_el('e')
    e.append(math_run(base))
    sub_node = m_el('sub')
    sub_node.append(math_run(sub))
    node.extend([e, sub_node])
    return node


def math_sup(base, sup):
    node = m_el('sSup')
    e = m_el('e')
    e.append(math_run(base))
    sup_node = m_el('sup')
    sup_node.append(math_run(sup))
    node.extend([e, sup_node])
    return node


def replace_regional_equation(paragraph):
    clear_paragraph(paragraph)
    omath_para = m_el('oMathPara')
    omath = m_el('oMath')
    omath.extend([
        math_sub('z', 'i'),
        math_run(' = '),
        math_sub('P', 'θ'),
        math_run('('),
        math_run('E'),
        math_run('('),
        math_sub('r', 'i'),
        math_run(')) ∈ '),
        math_sup('R', '128'),
        math_run(',     i ∈ {1,…,K}.'),
    ])
    omath_para.append(omath)
    paragraph._p.append(omath_para)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    p_rep = next(p for p in doc.paragraphs if p.text.startswith('Each retinal box was divided'))
    set_paragraph_text(
        p_rep,
        'Each retinal foreground box was partitioned into K=9 rectangular regions. The full-view, center-plus-cross, and representative regional-crop examples in Figure 2 were generated from the same source UWF image; black tiles denote regions excluded by the visibility mask. Each crop was resized to 256×256, center-cropped to 224×224, and normalized using the ImageNet channel-wise mean and standard deviation. For each regional crop, a local RetinaRadar EfficientNet-B0 checkpoint generated a frozen global pooled feature vector. The checkpoint contained 17 retinal image-quality outputs covering laterality, image type, artifacts, clarity, illumination, contrast, field, and usability. It was used here solely as a frozen retinal image-quality encoder. The projection layer, visibility-conditioned attention module, and DR classifier were trained on UWF-DR. The shared feature path is shown in Figure 2B. The checkpoint identifier and checksum are provided in the accompanying code-and-results archive and Reference 23. [18,23]'
    )

    equation_index = next(
        i for i, p in enumerate(doc.paragraphs)
        if p._p.xpath('.//m:oMathPara') and 'zi' in ''.join(p._p.xpath('.//m:t/text()'))
    )
    replace_regional_equation(doc.paragraphs[equation_index])

    p_explain = next(p for p in doc.paragraphs if p.text.startswith('Here, r'))
    set_paragraph_text(
        p_explain,
        'Here, rᵢ denotes the i-th regional crop, E is the frozen EfficientNet-B0 encoder, and Pθ is a trainable projection module consisting of a linear layer, layer normalization, and ReLU activation. The resulting 128-dimensional vector zᵢ serves as the regional token.'
    )

    doc.save(OUTPUT)
    print(f'Wrote {OUTPUT}')


if __name__ == '__main__':
    main()
