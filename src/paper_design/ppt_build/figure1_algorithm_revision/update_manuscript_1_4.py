#!/usr/bin/env python3
"""Synchronize revised Figure 1, equations, and terminology into RegionSuff 1.4."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path("/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.3.docx")
OUTPUT = Path("/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.4.docx")
FIGURE = Path(
    "/Users/tongyue/Documents/77f /7hao/paper_design/ppt_build/"
    "figure1_algorithm_revision/figure1-revised-slide-01.png"
)


def set_run_font(run, *, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Times New Roman")
    rpr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rpr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def replace_body_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
        set_run_font(paragraph.runs[0], size=12)
    else:
        set_run_font(paragraph.add_run(text), size=12)


def replace_caption(paragraph, number, text):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    set_run_font(paragraph.add_run(f"Figure {number}. "), size=10.5, bold=True)
    set_run_font(paragraph.add_run(text), size=10.5)


def m_el(tag):
    return OxmlElement(f"m:{tag}")


def mr(text):
    r = m_el("r")
    rpr = m_el("rPr")
    sty = m_el("sty")
    sty.set(qn("m:val"), "p")
    rpr.append(sty)
    r.append(rpr)
    t = m_el("t")
    t.text = text
    r.append(t)
    return r


def sub(base_text, sub_text):
    node = m_el("sSub")
    pr = m_el("sSubPr")
    pr.append(m_el("ctrlPr"))
    node.append(pr)
    base = m_el("e")
    base.append(deepcopy(base_text) if hasattr(base_text, "tag") else mr(base_text))
    node.append(base)
    subnode = m_el("sub")
    subnode.append(deepcopy(sub_text) if hasattr(sub_text, "tag") else mr(sub_text))
    node.append(subnode)
    return node


def sup(base_text, sup_text):
    node = m_el("sSup")
    pr = m_el("sSupPr")
    pr.append(m_el("ctrlPr"))
    node.append(pr)
    base = m_el("e")
    base.append(deepcopy(base_text) if hasattr(base_text, "tag") else mr(base_text))
    node.append(base)
    supnode = m_el("sup")
    supnode.append(deepcopy(sup_text) if hasattr(sup_text, "tag") else mr(sup_text))
    node.append(supnode)
    return node


def radical(nodes):
    node = m_el("rad")
    pr = m_el("radPr")
    hide = m_el("degHide")
    hide.set(qn("m:val"), "1")
    pr.append(hide)
    node.append(pr)
    degree = m_el("deg")
    node.append(degree)
    expr = m_el("e")
    for item in nodes:
        expr.append(item)
    node.append(expr)
    return node


def frac(num_nodes, den_nodes):
    node = m_el("f")
    node.append(m_el("fPr"))
    numerator = m_el("num")
    for item in num_nodes:
        numerator.append(item)
    node.append(numerator)
    denominator = m_el("den")
    for item in den_nodes:
        denominator.append(item)
    node.append(denominator)
    return node


def aperture_equation_paragraph(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    omath_para = m_el("oMathPara")
    omath = m_el("oMath")
    nodes = [
        sub("A", "rectangle"),
        mr(" = "),
        sup("f", "2"),
        mr(",      "),
        sub("r", "circle"),
        mr(" = "),
        frac([mr("f")], [radical([mr("π")])]),
        mr(",      "),
        sub("A", "circle"),
        mr(" = π"),
        sup(sub("r", "circle"), "2"),
        mr(" = "),
        sup("f", "2"),
    ]
    for node in nodes:
        omath.append(node)
    omath_para.append(omath)
    paragraph._p.append(omath_para)
    return paragraph


def replace_figure_blob(doc, paragraph, image_bytes):
    embeds = paragraph._p.xpath(".//a:blip/@r:embed")
    if len(embeds) != 1:
        raise RuntimeError(f"Expected one Figure 1 image relationship, found {embeds}")
    image_part = doc.part.related_parts[embeds[0]]
    image_part._blob = image_bytes
    alt = (
        "RegionSuff study overview: hidden evidence gap, controlled regional and pixel-space "
        "evidence loss, visibility-conditioned attention equations, primary DR severity output, "
        "and downstream paired-view underestimation audit."
    )
    for docpr in paragraph._p.xpath(".//wp:docPr"):
        docpr.set("name", "Figure 1")
        docpr.set("descr", alt)


def find_paragraph(doc, startswith):
    matches = [p for p in doc.paragraphs if p.text.startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {startswith!r}, found {len(matches)}")
    return matches[0]


def main():
    if not SOURCE.exists() or not FIGURE.exists():
        raise FileNotFoundError((SOURCE, FIGURE))
    doc = Document(SOURCE)

    figure_paragraphs = [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]
    if not figure_paragraphs:
        raise RuntimeError("No inline figures found")
    replace_figure_blob(doc, figure_paragraphs[0], FIGURE.read_bytes())

    replace_caption(
        find_paragraph(doc, "Figure 1. "),
        1,
        "Study overview. (A) Technical gradability and task-relevant evidence availability are distinct: a clear image may still omit evidence required for DR grading. (B) Region masks alter which 3×3 tokens remain visible, while centered rectangular and area-matched circular pixel-space apertures alter area, location, and shape on the same source image without changing its DR label. (C) RegionSuff maps each region rᵢ to a token zᵢ, computes evidence scores eᵢ, normalizes attention aᵢ(m) only across visible tokens, aggregates h(m), and produces the primary DR probability vector p(m) and class prediction ŷ(m). (D) Internal robustness, one-region intervention, and independent UWF replication evaluate the severity model; the downstream paired-view audit first obtains p_F and p_L from full and limited masks and then estimates underestimation risk q(g). All restrictions are synthetic; no clinician-adjudicated sufficiency labels or deployable single-view risk rule are provided.",
    )

    replacements = {
        "The analyses were designed to characterize": (
            "The analyses were designed to characterize how the evidence available to the DR classifier changed as retinal coverage was restricted. The masks therefore represent controlled, synthetic restrictions of the observed retina rather than real pairs of conventional-field and UWF acquisitions. Because the datasets do not provide clinician-adjudicated labels of diagnostic sufficiency, the study focuses on model-specific evidence availability rather than clinical sufficiency itself. For the paired-view audit, the same fitted RegionSuff severity model was evaluated under a full mask m_F and a limited mask m_L to obtain p_F and p_L; a separate downstream risk model then analyzed their paired differences. This module is an analytical comparison and does not operate on a single image."
        ),
        "Pixel-level apertures were applied": (
            "Pixel-space apertures were applied before feature extraction while the source image, reference DR label, retinal foreground box, and 3×3 feature coordinates remained fixed. Centered rectangles and area-matched circles used linear fractions f of 0.35, 0.45, 0.55, 0.65, 0.75, and 0.90. A centered rectangle with f=0.55 was also translated left, right, superiorly, or inferiorly to isolate location, whereas the centered rectangle with f=0.35 served as the severe rectangular aperture. For a rectangle, retained width and height were f times the retinal box. For the normalized circle, radius r_circle=f/√π matched the nominal aperture area f² of the rectangle. Actual retained retinal tissue could still differ because of retinal shape, image borders, and black background."
        ),
        "The attention network applies": (
            "The attention network applies a 128→48 linear layer, tanh activation, and a scalar output to each regional token. It first computes the evidence score eᵢ for every region. In the implementation, logits for mᵢ=0 were set to −10⁴ before softmax; the equations below give the mathematically equivalent visible-set form for a nonempty visibility vector m."
        ),
        "The binary vector m identifies visible regions": (
            "The binary visibility vector m identifies the evidence available for the current view, with mᵢ=1 for a visible region and mᵢ=0 otherwise. Consequently, aᵢ(m)=0 for an unavailable region and the attention weights sum to one only over visible tokens. Their weighted sum forms h(m), and the classifier produces p(m) over Normal, NPDR, and PDR. The primary severity output is ŷ(m)=arg max_c p_c(m). The unweighted cross-entropy objective was used because the final implementation did not apply class weights."
        ),
        "For each image, the paired audit compared": (
            "For each image, the same fitted RegionSuff severity model was evaluated twice: the full visibility mask m_F produced probabilities p_F, and the center-plus-cross limited mask m_L produced p_L. These two outputs were passed to a downstream paired-view audit rather than treated as parallel outputs of the regional attention module. The resulting 18-dimensional feature contains four probability blocks (12 values), three expected-severity terms, two predictive entropies, and one symmetric Kullback–Leibler divergence."
        ),
        "Here, ε=10⁻⁸ was used": (
            "Here, ε=10⁻⁸ was used for numerical stability; s_F and s_L are the expected severity scores derived from p_F and p_L; H_F and H_L are their predictive entropies; and D_SKL is the average of the two directed KL divergences. The standardization moments μ_g and σ_g were estimated only from the risk-development partition. The logistic risk model used L2 regularization (C=0.1), balanced class weights, and liblinear optimization. In the fixed-split analysis, the same validation predictions were used to fit the logistic model and choose the risk threshold; the held-out test ranking is informative, but the threshold claim is exploratory because model fitting and calibration were not separated."
        ),
        "RegionSuff was developed using 1,630 UWF-DR records": (
            "RegionSuff was developed using 1,630 UWF-DR records with group-disjoint model partitions and five training seeds. [21] The primary analysis compared full-view, center-only, and center-plus-cross inference, while additional experiments evaluated regional granularity, masking probability, aggregation strategy, encoder choice, continuous pixel apertures, and one-region occlusion. The occlusion analysis examined whether learned attention was functionally aligned with the performance loss caused by removing individual regions. External evaluation used the official 2,597-image MMRDR-UWF test set without target-domain adaptation or model selection. [22] A downstream paired-view risk model, applied only after full and center-plus-cross probability vectors were available, examined limited-view underestimation and the transportability of risk thresholds under grouped nested evaluation. Together, these analyses assess whether RegionSuff can characterize model behavior as the amount and spatial distribution of visible retinal evidence change. Because the view restrictions are synthetic and the datasets do not contain clinician-adjudicated sufficiency labels, RegionSuff is evaluated as a model-specific evidence-auditing framework rather than as a clinical test of diagnostic sufficiency."
        ),
        "The project archive contains": (
            "The project archive contains the analysis scripts, frozen experiment outputs, paired-bootstrap results, validation-list audit, manuscript builder, and PNG/PDF/SVG figure sources. Original Exp01–Exp60 outputs were not overwritten while preparing version 1.4."
        ),
    }
    for prefix, text in replacements.items():
        replace_body_paragraph(find_paragraph(doc, prefix), text)

    aperture_paragraph = find_paragraph(doc, "Pixel-space apertures were applied")
    eq_paragraph = aperture_equation_paragraph(doc)
    aperture_paragraph._p.addnext(eq_paragraph._p)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
