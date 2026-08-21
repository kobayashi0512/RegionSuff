from __future__ import annotations

import argparse
import shutil
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_run_font(run, size: float, bold: bool | None = None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Times New Roman")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")
        borders.append(node)
    for edge in ("left", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def set_cell_margins(cell, margin_twips=40):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def style_tables_and_captions(doc):
    for table in doc.tables:
        table.style = "Normal Table"
        set_table_borders(table)
        compact = len(table.columns) >= 8
        table_font_size = 9.0 if compact else 10.5
        header = table.rows[0]
        header_pr = header._tr.get_or_add_trPr()
        if header_pr.find(qn("w:tblHeader")) is None:
            header_pr.append(OxmlElement("w:tblHeader"))
        for row_index, row in enumerate(table.rows):
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                if compact:
                    set_cell_margins(cell)
                tc_pr = cell._tc.get_or_add_tcPr()
                shading = tc_pr.find(qn("w:shd"))
                if shading is not None:
                    tc_pr.remove(shading)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if row_index == 0
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_run_font(run, table_font_size, bold=(row_index == 0))

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Table "):
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            for index, run in enumerate(paragraph.runs):
                set_run_font(run, 12, bold=(index == 0))


def move_table_captions_above_tables(doc):
    body = doc.element.body
    table_nodes = [child for child in body if child.tag == qn("w:tbl")]
    for table_node in table_nodes:
        sibling = table_node.getnext()
        caption_node = None
        while sibling is not None:
            if sibling.tag == qn("w:p"):
                text = "".join(sibling.itertext()).strip()
                if text.startswith("Table "):
                    caption_node = sibling
                    break
            sibling = sibling.getnext()
        if caption_node is None:
            continue
        body.remove(caption_node)
        body.insert(body.index(table_node), caption_node)


def move_pair_details_after_figure(doc):
    paragraphs = doc.paragraphs
    heading_index = next(
        i for i, p in enumerate(paragraphs)
        if p.text.strip() == "2.1 Study design, endpoints, and reporting boundary"
    )
    p2_index = heading_index + 2
    image_index = next(
        i for i, p in enumerate(paragraphs[p2_index + 1 :], p2_index + 1)
        if p._p.xpath(".//wp:inline")
    )
    caption_index = next(
        i for i, p in enumerate(paragraphs[image_index + 1 :], image_index + 1)
        if p.text.startswith("Figure 1.")
    )
    details = paragraphs[p2_index + 1 : image_index]
    if details:
        parent = details[0]._p.getparent()
        for paragraph in details:
            parent.remove(paragraph._p)
        anchor = paragraphs[caption_index]._p
        for paragraph in details:
            anchor.addnext(paragraph._p)
            anchor = paragraph._p

    image = paragraphs[image_index]
    image.paragraph_format.space_before = Pt(0)
    image.paragraph_format.space_after = Pt(0)

    heading_22 = next(
        p for p in doc.paragraphs
        if p.text.strip() == "2.2 Datasets, records, and label harmonization"
    )
    ppr = heading_22._p.get_or_add_pPr()
    page_break = ppr.find(qn("w:pageBreakBefore"))
    if page_break is not None:
        ppr.remove(page_break)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if not args.source.is_file():
        raise FileNotFoundError(args.source)
    if args.output.resolve() == args.source.resolve():
        raise RuntimeError("The output must be a new DOCX.")

    doc = Document(args.source)
    move_pair_details_after_figure(doc)
    move_table_captions_above_tables(doc)
    style_tables_and_captions(doc)

    with tempfile.TemporaryDirectory(prefix="regionsuff_2_2_") as temp_dir:
        staged = Path(temp_dir) / args.output.name
        doc.save(staged)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(staged, args.output)
    print(f"Styled tables and reflowed Figure 1 in {args.output}")


if __name__ == "__main__":
    main()
