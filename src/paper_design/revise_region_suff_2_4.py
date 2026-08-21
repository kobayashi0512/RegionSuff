from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.3.docx')
OUTPUT = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.4.docx')


def clear_paragraph(paragraph):
    p = paragraph._p
    p_pr = p.pPr
    for child in list(p):
        if child is not p_pr:
            p.remove(child)


def set_caption(paragraph, text, prefix):
    clear_paragraph(paragraph)
    run = paragraph.add_run(prefix)
    run.bold = True
    paragraph.add_run(text[len(prefix):])


def set_run_font(run, size=10.5, bold=None):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{key}'), 'Times New Roman')


def set_cell_text(cell, text, header=False):
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, 10.5, bold=header)


def main():
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    table_caption = next(p for p in doc.paragraphs if p.text.startswith('Table 1.'))
    set_caption(
        table_caption,
        'Table 1. Datasets, study records, model partitions, and experimental roles.',
        'Table 1.'
    )

    table = doc.tables[0]
    headers = [
        'Dataset',
        'Records used in this study',
        'Partition or evaluation set',
        'Labels used',
        'Experimental role',
    ]
    row_values = [
        [
            'UWF-DR',
            '1,630',
            'Source training 876; validation 364; test 390',
            'Normal 496; NPDR 634; PDR 500',
            'Model development, controlled perturbation analyses, and internal evaluation',
        ],
        [
            'MMRDR-UWF',
            '2,597',
            'Official UWF external test subset',
            'Grade 0→Normal; grades 1–3→NPDR; grade 4→PDR',
            'Independent same-modality external evaluation; no target-domain adaptation',
        ],
    ]
    for ci, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[ci], text, header=True)
    for ri, values in enumerate(row_values, start=1):
        for ci, text in enumerate(values):
            set_cell_text(table.rows[ri].cells[ci], text, header=False)

    p_uwf = next(p for p in doc.paragraphs if p.text.startswith('UWF-DR provides'))
    p_uwf.text = (
        'UWF-DR contributed 1,630 publicly released 512×512 UWF image records acquired with an Optomap Daytona system and labeled as Normal, NPDR, or PDR. '
        'Because the release did not contain a dedicated patient-identifier column, we derived a grouping key from the first two underscore-separated filename fields. '
        'This procedure yielded 783 non-overlapping groups and source partitions of 876, 364, and 390 records for training, validation, and testing, respectively. [21]'
    )

    p_audit = next(p for p in doc.paragraphs if p.text.startswith('Audit of the released'))
    p_audit.text = (
        'An audit of the released metadata identified 1,608 unique filenames among the 1,630 records. Eleven filenames were duplicated, and two duplicated filenames were associated with conflicting labels. '
        'Group membership remained disjoint, and the 390-record test partition contained no repeated filename. A separate sensitivity analysis was performed on the validation list by retaining one row per filename and resolving the two conflicts with majority voting and a severe-class tie-break. '
        'The test set was not modified. Full-view, center-only, and center-plus-cross macro-F1 values in this analysis were 0.833, 0.790, and 0.832, respectively, closely matching the primary analysis.'
    )

    p_mm = next(p for p in doc.paragraphs if p.text.startswith('MMRDR includes'))
    p_mm.text = (
        'MMRDR contains CFP, OCT, and UWF modalities. We used only the official 2,597-image UWF test subset for independent same-modality external evaluation. '
        'Its original five-level DR labels were mapped to the prespecified three-class task: grade 0 to Normal, grades 1–3 to NPDR, and grade 4 to PDR. '
        'The resulting class counts were 800, 1,411, and 386. MMRDR images and labels were not used for source training, early stopping, calibration, model selection, or threshold selection. [22]'
    )

    doc.save(OUTPUT)
    print(f'Wrote {OUTPUT}')


if __name__ == '__main__':
    main()
