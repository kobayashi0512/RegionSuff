from pathlib import Path
import shutil

from docx import Document


SOURCE = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.6.docx')
OUTPUT = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.7.docx')


def main():
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    p_intro = next(p for p in doc.paragraphs if p.text.startswith('The attention network applies'))
    p_intro.text = (
        'The attention module applies a 128-to-48 linear layer, followed by a tanh activation and a 48-to-1 linear readout, to each regional token. '
        'It produces a regional attention score eᵢ for every region. In the implementation, logits corresponding to unavailable regions (mᵢ=0) were set to −10⁴ before softmax. '
        'For a nonempty visibility vector, this is equivalent to the following visible-set formulation:'
    )

    p_explain = next(p for p in doc.paragraphs if p.text.startswith('The binary visibility vector'))
    p_explain.text = (
        'The binary visibility vector m identifies the regions available in the current view, with mᵢ=1 for a visible region and mᵢ=0 otherwise. '
        'Thus, unavailable regions receive effectively zero attention, and the weights are normalized only over visible tokens. Their weighted sum forms h(m), from which the classifier produces probabilities for Normal, NPDR, and PDR. '
        'The predicted class is ŷ(m)=arg max_c p_c(m), with c∈{0,1,2}. The all-empty mask case was handled using the fallback rule described in Section 2.6. '
        'The model was optimized with unweighted cross-entropy because class weights were not used in the locked implementation. Here, B is the minibatch size, yₙ is the reference class for sample n, and mₙ is its visibility mask.'
    )

    doc.save(OUTPUT)
    print(f'Wrote {OUTPUT}')


if __name__ == '__main__':
    main()
