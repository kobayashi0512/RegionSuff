from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.4.docx')
OUTPUT = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_2.5.docx')


def clear_paragraph(paragraph):
    p = paragraph._p
    p_pr = p.pPr
    for child in list(p):
        if child is not p_pr:
            p.remove(child)


def set_paragraph_text(paragraph, text):
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def m_el(tag):
    return OxmlElement(f'm:{tag}')


def math_run(text):
    run = m_el('r')
    text_node = m_el('t')
    text_node.text = text
    run.append(text_node)
    return run


def math_sub(base, sub):
    node = m_el('sSub')
    e = m_el('e')
    e.append(math_run(base))
    sub_node = m_el('sub')
    sub_node.append(math_run(sub))
    node.extend([e, sub_node])
    return node


def math_sup(base, sup):
    node = m_el('sSup')
    e = m_el('e')
    e.append(math_run(base))
    sup_node = m_el('sup')
    sup_node.append(math_run(sup))
    node.extend([e, sup_node])
    return node


def replace_aperture_equation(paragraph):
    clear_paragraph(paragraph)
    omath_para = m_el('oMathPara')
    omath = m_el('oMath')
    omath.extend([
        math_sub('A', 'rectangle'),
        math_run(' / '),
        math_sub('A', 'box'),
        math_run(' = '),
        math_sup('f', '2'),
        math_run(',     '),
        math_sub('r', 'circle'),
        math_run('(norm) = f / √π,     '),
        math_sub('A', 'circle'),
        math_run(' / '),
        math_sub('A', 'box'),
        math_run(' = '),
        math_sup('f', '2'),
        math_run('.'),
    ])
    omath_para.append(omath)
    paragraph._p.append(omath_para)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    p_foreground = next(p for p in doc.paragraphs if p.text.startswith('After pixel intensities'))
    set_paragraph_text(
        p_foreground,
        'After scaling each image channel to [0,1], we computed the per-pixel mean intensity and defined the retinal foreground as pixels whose mean intensity exceeded 0.035. The foreground box was then obtained as the tight axis-aligned bounding rectangle of these pixels. If the detected foreground occupied less than 8% of the image frame, the full frame was used as a fallback. Region-level views were generated from a 1-indexed 3×3 row-major grid: full-view retained all nine regions, center-only retained region 5, and center-plus-cross retained regions 2, 4, 5, 6, and 8.'
    )

    p_aperture = next(p for p in doc.paragraphs if p.text.startswith('Pixel-space apertures'))
    set_paragraph_text(
        p_aperture,
        'Pixel-space apertures were applied before feature extraction. The source image, reference DR label, retinal foreground box, and 3×3 regional coordinates were kept fixed across all interventions. The parameter f denoted the retained width and height as fractions of the foreground box. Centered rectangles and nominally area-matched circles were evaluated at f=0.35, 0.45, 0.55, 0.65, 0.75, and 0.90. A rectangle with f=0.55 was additionally translated leftward, rightward, superiorly, and inferiorly to examine location effects, whereas the centered rectangle with f=0.35 served as the severe rectangular aperture.'
    )

    equation_index = next(
        i for i, p in enumerate(doc.paragraphs)
        if p._p.xpath('.//m:oMathPara') and 'Arectangle' in ''.join(p._p.xpath('.//m:t/text()'))
    )
    equation_paragraph = doc.paragraphs[equation_index]
    replace_aperture_equation(equation_paragraph)

    # Add the dimensional qualification immediately after the equation.
    following = doc.paragraphs[equation_index + 1]
    if not following.text.startswith('Thus, the rectangle'):
        following.insert_paragraph_before(
            'Thus, the rectangle and circle were matched by nominal aperture area rather than by the exact amount of retained retinal tissue. The actual retained tissue could still differ because of retinal shape, image boundaries, and the black background.',
            style=p_aperture.style.name,
        )

    doc.save(OUTPUT)
    print(f'Wrote {OUTPUT}')


if __name__ == '__main__':
    main()
