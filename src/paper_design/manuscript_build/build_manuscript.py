#!/usr/bin/env python3
from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path("/Users/tongyue/Documents/77f /7hao")
REFERENCE = Path("/Users/tongyue/Documents/ppf/第二论文/VOCT_2.0.docx")
FIG_DIR = ROOT / "paper_design/figures"
OUTPUT = ROOT / "RegionSufficiency_论文稿件_已插入框架与算法图.docx"

INK = "3F3027"
BROWN = "704C2B"
UMBER = "965E4A"
ORANGE = "E49F5D"
OLIVE = "8B9A7C"
ROSE = "C74282"
PAPER = "FFFDFC"
PALE = "F7ECE6"
PALE_OLIVE = "E5E8D9"


def set_run_font(run, name="Times New Roman", size=11, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        tag = "w:" + edge
        node = tc_borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_borders.append(node)
        for key, value in edges[edge].items():
            node.set(qn("w:" + key), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def clear_document_body(doc: Document):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True, page_break_before=False):
    p_pr = paragraph._p.get_or_add_pPr()
    for tag, enabled in (("keepNext", keep_next), ("keepLines", keep_lines), ("pageBreakBefore", page_break_before)):
        old = p_pr.find(qn(f"w:{tag}"))
        if old is not None:
            p_pr.remove(old)
        if enabled:
            p_pr.append(OxmlElement(f"w:{tag}"))


def add_bottom_rule(paragraph, color=ORANGE, size=10, space=3):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction, display=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=UMBER)
    return run


def style_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.12
    pf.space_after = Pt(5)

    for name, size, color in (("Heading 1", 14, BROWN), ("Heading 2", 12, UMBER), ("Heading 3", 11, OLIVE)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    if "Caption" in [s.name for s in doc.styles]:
        style = doc.styles["Caption"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(9)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.0


def add_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = "REGIONSUFFICIENCY  ·  MANUSCRIPT DRAFT"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        set_run_font(run, name="Arial", size=8.5, bold=True, color=UMBER)
    add_bottom_rule(p, color="E8C6AA", size=6, space=2)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RegionSufficiency  ·  ")
    set_run_font(r, name="Arial", size=8.5, color=UMBER)
    add_field(p, "PAGE", "1")


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("When Is a Fundus Image Diagnostically Sufficient?")
    set_run_font(r, name="Arial", size=18, bold=True, color=BROWN)
    r.add_break()
    r2 = p.add_run("Visibility-Conditioned Regional Evidence Aggregation under Limited Fields of View")
    set_run_font(r2, name="Arial", size=15, bold=True, color=UMBER)
    add_bottom_rule(p, color=ORANGE, size=12, space=8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Figure-integrated research manuscript draft · RegionSufficiency study")
    set_run_font(r, size=9.5, italic=True, color=OLIVE)


def add_heading(doc: Document, text: str, level=1, page_break=False):
    p = doc.add_heading(text, level=level)
    set_paragraph_keep(p, keep_next=True, page_break_before=page_break)
    if level == 1:
        add_bottom_rule(p, color="E8C6AA", size=6, space=2)
    return p


def add_body(doc: Document, text: str, first_line=True, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.65) if first_line else None
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    set_paragraph_keep(p, keep_next=keep_next, keep_lines=True)
    return p


def add_key_value(doc: Document, label: str, value: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + " ")
    set_run_font(r, size=10.5, bold=True, color=UMBER)
    r = p.add_run(value)
    set_run_font(r, size=10.5, color=INK)
    return p


def add_equation(doc: Document, equation: str, explanation: str):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(8.7)
    table.columns[1].width = Cm(7.5)
    left, right = table.rows[0].cells
    left.width, right.width = Cm(8.7), Cm(7.5)
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell, top={"val": "single", "sz": "4", "color": "E8C6AA"}, bottom={"val": "single", "sz": "4", "color": "E8C6AA"}, left={"val": "nil"}, right={"val": "nil"})
    set_cell_shading(left, PALE)
    set_cell_shading(right, PAPER)
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(equation)
    set_run_font(r, name="Cambria Math", size=10.5, italic=True, color=BROWN)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(explanation)
    set_run_font(r, size=9.2, color=INK)
    table.rows[0].height = Cm(1.0)
    set_cant_split(table.rows[0])
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc: Document, image_path: Path, number: int, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    set_paragraph_keep(p, keep_next=True, keep_lines=True)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(6.55))

    p = doc.add_paragraph(style="Caption" if "Caption" in [s.name for s in doc.styles] else None)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"Figure {number}. ")
    set_run_font(r, size=9, bold=True, color=UMBER)
    r = p.add_run(caption)
    set_run_font(r, size=9, color=INK)
    set_paragraph_keep(p, keep_next=False, keep_lines=True)
    return p


def add_dataset_table(doc: Document):
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.2), Cm(2.0), Cm(3.0), Cm(3.9), Cm(4.2)]
    headers = ["Dataset", "Images", "Modality", "Role", "Leakage control"]
    for i, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        set_cell_shading(cell, BROWN)
        set_cell_border(cell, bottom={"val": "single", "sz": "6", "color": ORANGE})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, name="Arial", size=8.5, bold=True, color="FFFFFF")
    set_repeat_table_header(table.rows[0])
    rows = [
        ("UWF-DR", "1,630", "UWF fundus", "Development and controlled-view stress tests", "Patient-grouped train/validation/test; five random seeds"),
        ("MMRDR-UWF", "2,597", "UWF fundus", "Independent same-modality external replication", "Official patient-level test split; no target labels for training or selection"),
    ]
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, text, width) in enumerate(zip(cells, row, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, PAPER if ridx % 2 == 0 else PALE)
            set_cell_border(cell, bottom={"val": "single", "sz": "3", "color": "E8C6AA"})
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in (0, 3, 4) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=8.5, bold=(i == 0), color=INK)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run("Table 1. ")
    set_run_font(r, size=9, bold=True, color=UMBER)
    r = cap.add_run("Primary datasets and their non-overlapping roles in model development and external evaluation.")
    set_run_font(r, size=9, color=INK)
    return table


def build():
    doc = Document(str(REFERENCE))
    clear_document_body(doc)
    style_document(doc)
    add_header_footer(doc)
    add_title(doc)

    add_heading(doc, "Abstract", 1)
    add_body(doc, "Automated fundus-image analysis usually assumes that an acquired image already contains enough disease-relevant evidence. This assumption can fail even when the image is sharp and apparently usable, because diagnostically informative retinal regions may be absent from the field of view. We therefore formulate diagnostic sufficiency as a visibility-conditioned inference problem rather than a conventional image-quality problem.", first_line=False)
    add_body(doc, "RegionSufficiency partitions the retinal field into a 3×3 region bank, extracts frozen fundus-specific deep features, and aggregates only visible regional tokens with a masked-attention model trained under stochastic region removal. The model jointly supports diabetic-retinopathy severity estimation and a paired-view audit of limited-view underestimation risk. Development used 1,630 UWF-DR images with patient-grouped splits and five random seeds. Independent same-modality replication used the official 2,597-image MMRDR-UWF test set without target-label training, calibration, early stopping, or model selection.", first_line=False)
    add_body(doc, "On internal testing, macro-F1 was 0.830±0.017 under the full view, 0.790±0.012 under the center view, and 0.829±0.017 under the center-plus-cross view. On MMRDR-UWF, the corresponding macro-F1 values were 0.625±0.010, 0.600±0.004, and 0.642±0.005. These non-monotonic view effects indicate that the spatial composition of visible evidence matters beyond visible area alone. The current risk component is a paired-view research audit and is not presented as a clinically validated single-view safety system.", first_line=False)
    add_key_value(doc, "Keywords:", "diagnostic sufficiency; fundus imaging; limited field of view; regional evidence; masked attention; selective prediction; diabetic retinopathy")

    add_heading(doc, "1. Introduction", 1, page_break=True)
    add_body(doc, "Automated diabetic-retinopathy grading has the potential to extend large-scale retinal screening, but real acquisition conditions frequently depart from curated benchmark images. A photograph can be affected by incomplete retinal coverage, black borders, eyelashes, glare, uneven illumination, device-specific geometry, or operator-dependent centering. Existing pipelines often treat these problems as generic image-quality failures before disease inference begins.")
    add_body(doc, "That workflow hides a stronger assumption: once an image passes a quality gate, it is presumed to contain sufficient evidence for the target diagnosis. Yet visual clarity and diagnostic sufficiency answer different questions. Image quality concerns sharpness, exposure, contrast, and artifacts. Diagnostic sufficiency concerns whether the currently visible retinal regions contain the evidence required for a reliable disease judgment. A clear central image can therefore remain diagnostically incomplete, while a nonstandard but evidence-rich view may still support the task.")
    add_body(doc, "Conventional full-image classifiers do not resolve this distinction because they are forced to return a class for every input. Generic uncertainty measures quantify predictive ambiguity but do not explicitly encode which retinal regions were unavailable. Likewise, global pooling can down-weight a visible region, but it cannot distinguish a region assigned little importance from a region that was never observed. This conflation is particularly problematic when peripheral lesions or distributed vascular changes contribute to disease severity.")
    add_body(doc, "We introduce RegionSufficiency, a visibility-conditioned regional evidence framework that represents the retina as a set of spatial evidence units and explicitly masks unavailable units before aggregation. During training, random region removal teaches the model to reason under incomplete views. During evaluation, controlled masks vary visible area, location, shape, and region composition without changing the patient or disease label. The resulting study design links the conceptual gap, the regional algorithm, the internal stress tests, the independent UWF replication, and the selective risk audit in one falsifiable workflow (Figure 1).", keep_next=True)

    add_figure(
        doc,
        FIG_DIR / "Figure1_Study_Overview_RMB20.png",
        1,
        "Study overview. (A) Traditional image-quality assessment may label a clear image as acceptable even when disease-relevant retinal regions were not captured. (B) The model was developed on 1,630 UWF-DR images with patient-grouped splits and evaluated on the independent 2,597-image MMRDR-UWF official test set. Controlled views varied visible area, location, shape, and region composition. (C) RegionSufficiency encodes a 3×3 retinal region bank and aggregates only visible regional evidence. (D) Evaluation included controlled evidence topology, internal multi-seed performance, independent UWF replication, and selective risk analysis. The framework supports acceptance-or-deferral research but is not presented as a clinically validated safety system."
    )

    add_body(doc, "The study asks four linked questions: whether diagnostic sufficiency is only a function of visible area; why visibility-conditioned regional aggregation is useful; whether the evidence pattern replicates in an independent UWF dataset; and whether limited-view underestimation risk can be ranked to support selective deferral. Our contribution is therefore not a claim that one crop is universally superior, but a task definition and evaluation protocol for identifying when the observed retinal evidence is insufficient for a specified model and disease target.")

    add_heading(doc, "2. Materials and Methods", 1, page_break=True)
    add_heading(doc, "2.1 Study design and datasets", 2)
    add_body(doc, "The study separated model development from external evaluation. UWF-DR was used for patient-grouped training, validation, internal testing, controlled-view construction, and ablation analysis. MMRDR-UWF was reserved for independent same-modality replication. Its labels did not enter representation learning, classifier training, threshold calibration, early stopping, or model selection.")
    add_dataset_table(doc)

    add_heading(doc, "2.2 Patient grouping and leakage control", 2)
    add_body(doc, "All development splits were grouped by patient identity so that images from the same patient could not appear across training, validation, and test subsets. The final configuration was evaluated across five random seeds. Feature normalization was fitted on training data only; model selection used validation macro-F1 under the prespecified limited-view mask; the held-out internal test set was used once for final reporting. External MMRDR-UWF results were computed using source-trained models without target-label adaptation.")

    add_heading(doc, "2.3 Operational definition of diagnostic sufficiency", 2)
    add_body(doc, "Let x denote a fundus image, y its disease-severity label, and m∈{0,1}⁹ a visibility mask over nine retinal regions. Diagnostic sufficiency is operationalized through the behavior of the same severity model when its observable evidence changes from a full mask mF to a limited mask mL. The primary endpoint is macro-F1 under the center-plus-cross limited view. Secondary safety-oriented endpoints include severity mean absolute error, underestimation rate, and PDR recall. The underestimation event is I[argmax p(y|x,mL)<y].")

    add_heading(doc, "2.4 RegionSufficiency architecture", 2)
    add_body(doc, "The model has four stages: controlled view construction, frozen regional representation, visibility-conditioned evidence aggregation, and severity/risk decision. Figure 2 gives the complete data flow and separates the trainable disease model from the paired-view risk audit.", keep_next=True)

    add_equation(doc, "zᵢ = P(E(rᵢ)),  i = 1,…,9", "Each retinal region rᵢ is encoded by the frozen RetinaRadar EfficientNet-B0 encoder E and mapped by trainable projection P to a regional token zᵢ.")
    add_equation(doc, "aᵢ(m) = mᵢ exp{s(zᵢ)} / Σⱼ mⱼ exp{s(zⱼ)}", "The binary visibility mask m removes unavailable regions before attention normalization. Invisible regions therefore receive exactly zero aggregation weight.")
    add_equation(doc, "h(m) = Σᵢ aᵢ(m)zᵢ;   p(m) = softmax(Wch(m)+bc)", "The visible regional tokens form image representation h(m), which is mapped to the three severity classes Normal, NPDR, and PDR.")
    add_equation(doc, "φ = [pF,pL,pF−pL,|pF−pL|,EF,EL,EF−EL,HF,HL,DSKL] ∈ ℝ¹⁸", "The paired-view audit summarizes full-versus-limited probability shifts, expected severity, entropy, and symmetric KL divergence in an 18-dimensional evidence-gap vector.")
    add_equation(doc, "q = σ(βᵀφ+b);   accept ⇔ q≤τα", "A validation-fitted logistic model estimates limited-view underestimation risk q. The threshold τα is selected so the one-sided Clopper–Pearson upper bound is no greater than target risk α.")

    add_figure(
        doc,
        FIG_DIR / "Figure2_Method_Architecture_RMB20.png",
        2,
        "RegionSufficiency architecture and evaluation protocol. (A) A retinal region of interest is divided into a 3×3 region bank; full, center, and center-plus-cross masks define visible regions, and 30% stochastic region masking is used during training. (B) A shared frozen RetinaRadar EfficientNet-B0 encoder and a trainable projection produce regional tokens. (C) The same masked-attention severity model is evaluated under full and limited views. (D) Full- and limited-view probabilities form an 18-dimensional paired evidence-gap vector used by a validation-fitted logistic model to estimate limited-view underestimation risk. A threshold selected with a one-sided Clopper–Pearson upper bound determines accept/defer decisions. The current risk module is a paired-view audit, not a single-view deployment head."
    )

    add_heading(doc, "2.5 Controlled visibility protocols", 2)
    add_body(doc, "Visibility was perturbed at two complementary scales. Region-level protocols used binary subsets of the 3×3 bank, including the full, center, and center-plus-cross views. Pixel-level protocols varied the continuous visible fraction, translated a fixed-area window across retinal locations, and compared shapes at matched visible area. Because every view was generated from the same underlying image, these protocols isolate evidence visibility while holding patient and ground-truth severity constant. They are geometric proxies and should not be interpreted as real same-eye acquisitions from different cameras.")

    add_heading(doc, "2.6 Training and model selection", 2)
    add_body(doc, "The final configuration used a frozen RetinaRadar EfficientNet-B0 regional encoder, a trainable projection, attention aggregation, a 3×3 grid, and a 0.30 training mask ratio. Model selection was prespecified as validation macro-F1 under the center-plus-cross mask. If validation scores were tied within 0.005, the smaller model and lower validation underestimation rate were preferred. Internal performance was summarized as mean±standard deviation across seeds 0–4.")

    add_heading(doc, "2.7 Paired-view underestimation risk and selective calibration", 2)
    add_body(doc, "The risk head was fitted only on validation-set features derived from paired full and limited predictions and evaluated on held-out data. Risk thresholds were selected by maximizing coverage subject to a one-sided 95% Clopper–Pearson upper bound on accepted underestimation risk. Nested group cross-validation was used as an additional threshold-transport audit. Because the current feature vector includes the full-view prediction, it evaluates evidence loss under controlled pairing and is not directly deployable when only a single limited-view image is available.")

    add_heading(doc, "2.8 Statistical analysis", 2)
    add_body(doc, "Internal model stability was reported across five random seeds. Within-image view comparisons used paired resampling. External MMRDR-UWF analysis used 2,000 paired image-level bootstrap resamples with the same target images evaluated under each view. Confidence intervals from this procedure quantify uncertainty for within-dataset view differences and do not represent cross-hospital clinical validation.")

    add_heading(doc, "3. Results", 1, page_break=True)
    add_heading(doc, "3.1 Is diagnostic sufficiency determined by visible area alone?", 2)
    add_body(doc, "The controlled-view experiments did not support a simple monotonic relation between visible area and performance. Views of equal size produced different outcomes when translated to different retinal locations, and the center-plus-cross topology preserved more disease evidence than a center-only view. These results motivate reporting location and topology together with area rather than using a single field-of-view percentage as a proxy for sufficiency.")

    add_heading(doc, "3.2 Does visibility-conditioned regional aggregation remain stable across patient splits?", 2)
    add_body(doc, "Across five patient-grouped random seeds, internal macro-F1 was 0.830±0.017 for the full view, 0.790±0.012 for the center view, and 0.829±0.017 for the center-plus-cross view. The corresponding underestimation rates were 0.070±0.007, 0.093±0.006, and 0.075±0.012. The center-plus-cross topology therefore recovered most full-view performance while exposing the model to a restricted evidence pattern.")

    add_heading(doc, "3.3 Does the evidence topology replicate in independent UWF data?", 2)
    add_body(doc, "Without MMRDR label training or model selection, external macro-F1 was 0.625±0.010 under the full view, 0.600±0.004 under the center view, and 0.642±0.005 under the center-plus-cross view. Center-plus-cross also reduced severity MAE from 0.400±0.013 under the full view to 0.378±0.006. The non-monotonic pattern is treated as evidence that spatial composition and device-domain effects interact; it is not a recommendation to crop clinical UWF images.")

    add_heading(doc, "3.4 Can evidence-loss risk support selective deferral?", 2)
    add_body(doc, "In the fixed held-out evaluation, the paired-view risk head achieved AUROC 0.740 and supported a 62.8% accepted coverage with 2.9% underestimation among accepted samples at the 5% target. Nested group calibration yielded a lower and more variable AUROC of 0.653±0.034, with 56.7%±27.4% coverage and 5.0%±3.3% accepted underestimation at the same target. Both estimates are reported because the nested result more directly reflects threshold transport across patient groups.")

    add_heading(doc, "4. Discussion", 1)
    add_body(doc, "The central finding is that diagnostic sufficiency is a property of visible evidence composition, not a synonym for photographic quality or visible area. RegionSufficiency makes missingness explicit by masking unavailable regions before aggregation, enabling controlled tests of where and how much retinal evidence is retained. The model is not proposed as the strongest general-purpose DR classifier; its value lies in converting an ignored acquisition assumption into a measurable model behavior.")
    add_body(doc, "The external MMRDR-UWF results reinforce the topology hypothesis but also caution against simplistic interpretation. A center-plus-cross mask outperformed the nominal full view in this external setting, which could reflect suppression of device-specific peripheral artifacts, differences in retinal framing, or a favorable evidence-to-noise balance. This finding should be investigated with foreground-crop controls and prospective same-eye acquisitions before any clinical acquisition recommendation is made.")
    add_body(doc, "The paired-view risk audit shows that evidence loss can be ranked, but nested calibration exposes substantial variability. A future deployable system would require a single-view risk head, real acquisition metadata, and prospective calibration with clinician-defined sufficiency outcomes. Attention weights are therefore interpreted as aggregation parameters rather than lesion-localization ground truth.")

    add_heading(doc, "5. Limitations", 1)
    add_body(doc, "The study has no direct clinician annotation of diagnostic sufficiency; controlled masks are geometric proxies rather than true same-eye limited-field acquisitions; the primary disease target is diabetic retinopathy; only one independent same-modality UWF dataset is available; and the current risk head depends on paired full/limited predictions. Clinical deployment, acquisition guidance, and safety claims are outside the present evidence boundary.")

    add_heading(doc, "6. Conclusion", 1)
    add_body(doc, "A fundus image can be visually clear yet diagnostically insufficient. By representing the retina as visibility-conditioned regional evidence, RegionSufficiency provides a reproducible way to test this hidden failure mode under controlled fields of view. Internal and independent UWF results indicate that evidence topology matters beyond visible area alone, while selective-risk analyses define both a promising direction and the calibration work still required for deployment.")

    add_heading(doc, "Data and code availability", 1)
    add_body(doc, "Dataset use follows the access conditions of the original public resources. The complete local experiment registry, model configurations, and statistical outputs are maintained with the study record. Any public release should preserve the external MMRDR-UWF protocol without target-label training or model selection.")

    doc.core_properties.title = "RegionSufficiency manuscript with integrated framework and method figures"
    doc.core_properties.subject = "Diagnostic sufficiency in fundus imaging"
    doc.core_properties.keywords = "fundus imaging, diagnostic sufficiency, limited field of view, regional attention"
    doc.core_properties.comments = "Figure-integrated manuscript draft generated from the project evidence registry."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
