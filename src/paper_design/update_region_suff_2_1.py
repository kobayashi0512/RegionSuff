from __future__ import annotations

import argparse
import shutil
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def m_el(tag: str):
    return OxmlElement(f"m:{tag}")


def math_run(text: str):
    run = m_el("r")
    t = m_el("t")
    t.text = text
    run.append(t)
    return run


def math_sub(base: str, sub: str):
    node = m_el("sSub")
    e = m_el("e")
    e.append(math_run(base))
    sub_node = m_el("sub")
    sub_node.append(math_run(sub))
    node.extend([e, sub_node])
    return node


def add_probability_equation(paragraph):
    """Insert the rendered Word equation corresponding to the LaTeX form.

    LaTeX source:
      p_{\mathrm{full}} = p\!\left(m_{\mathrm{full}}\right),\qquad
      p_{\mathrm{limited}} = p\!\left(m_{\mathrm{limited}}\right).
    """
    omath_para = m_el("oMathPara")
    omath = m_el("oMath")
    omath.extend(
        [
            math_sub("p", "full"),
            math_run(" = p("),
            math_sub("m", "full"),
            math_run(")     "),
            math_sub("p", "limited"),
            math_run(" = p("),
            math_sub("m", "limited"),
            math_run(")"),
        ]
    )
    omath_para.append(omath)
    paragraph._p.append(omath_para)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def insert_paragraph_before(anchor, text: str = "", style: str = "Normal (Web)"):
    paragraph = anchor.insert_paragraph_before(text)
    paragraph.style = style
    return paragraph


def set_run_font(run, size: float, bold: bool | None = None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Times New Roman")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")
        borders.append(node)
    for edge in ("left", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def style_tables_and_captions(doc):
    """Match the reference paper: bold headers, 10.5-pt body, open tables."""
    for table in doc.tables:
        table.style = "Normal Table"
        set_table_borders(table)
        header = table.rows[0]
        header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row_index, row in enumerate(table.rows):
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                shading = tc_pr.find(qn("w:shd"))
                if shading is not None:
                    tc_pr.remove(shading)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_run_font(run, 10.5, bold=(row_index == 0))

    # Table captions sit above the table in the reference manuscript.
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Table "):
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            for index, run in enumerate(paragraph.runs):
                set_run_font(run, 12, bold=(index == 0))


def move_before_pair_details(image_paragraph, caption, paragraphs_to_move):
    """Place Figure 1 after the study-design paragraph, before pairwise details."""
    parent = paragraphs_to_move[0]._p.getparent()
    for paragraph in paragraphs_to_move:
        parent.remove(paragraph._p)
    anchor = caption._p
    for paragraph in paragraphs_to_move:
        anchor.addnext(paragraph._p)
        anchor = paragraph._p

    heading = next(
        p for p in parent if p.tag == qn("w:p") and "2.2 Datasets" in "".join(p.itertext())
    )
    ppr = heading.find(qn("w:pPr"))
    if ppr is not None:
        page_break = ppr.find(qn("w:pageBreakBefore"))
        if page_break is not None:
            ppr.remove(page_break)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    if not args.source.is_file():
        raise FileNotFoundError(args.source)
    if args.output.resolve() == args.source.resolve():
        raise RuntimeError("The output must be a new DOCX; the source is preserved.")

    doc = Document(args.source)
    paragraphs = doc.paragraphs
    heading_index = next(
        i for i, p in enumerate(paragraphs)
        if p.text.strip() == "2.1 Study design, endpoints, and reporting boundary"
    )
    if not paragraphs[heading_index + 1].text.startswith("Using publicly available, de-identified data"):
        raise RuntimeError("The expected RegionSuff section structure was not found.")

    p1 = paragraphs[heading_index + 1]
    p2 = paragraphs[heading_index + 2]
    image_paragraph = paragraphs[heading_index + 3]
    caption = paragraphs[heading_index + 4]

    p1.text = (
        "Using publicly available, de-identified data, we conducted a retrospective computational study without additional image acquisition. "
        "Model performance was evaluated on the held-out test set under three predefined visibility settings: full-view, center-only, and center-plus-cross. "
        "Macro-F1 for three-class DR grading was defined as the primary endpoint. Accuracy, ordinal severity mean absolute error (MAE), and underestimation rate "
        "were included as secondary endpoints. Source-domain training was repeated across five random seeds to characterize variability arising from model optimization. "
        "For analyses with per-image predictions, differences between visibility settings were summarized using image-level paired bootstrap 95% confidence intervals. [28]"
    )

    p2.text = (
        "The study was designed to examine how the evidence available to a fixed DR classifier changes when retinal coverage is restricted. "
        "Region-level masks and pixel-space apertures were therefore treated as controlled interventions applied to the observed retina. They were not interpreted as real paired acquisitions "
        "from conventional-field and UWF imaging. In the absence of clinician-adjudicated labels for diagnostic sufficiency, the operational construct studied here was "
        "model-specific retinal evidence availability."
    )

    p3 = insert_paragraph_before(
        image_paragraph,
        "For the paired-view audit, the same fitted RegionSuff severity model was evaluated under a full-view mask m_full and a limited-view mask m_limited, yielding the corresponding probability vectors:",
    )
    equation = insert_paragraph_before(image_paragraph)
    add_probability_equation(equation)
    p4 = insert_paragraph_before(
        image_paragraph,
        "A separate downstream risk model then analyzed the paired differences between these two predictions. This analysis requires both visibility conditions and is intended to characterize prediction shifts under restricted views rather than to operate as an independent single-image triage rule.",
    )

    caption.text = (
        "Figure 1. RegionSuff study overview. (A) Technical gradability and task-relevant evidence availability are distinct: a clear image can omit evidence required for DR grading. "
        "(B) Region-level masks vary the visible 3×3 evidence tokens. Pixel-space apertures preserve a centered rectangle (f = 0.55), an area-matched circle (f = 0.55), or a severe rectangle (f = 0.35) from the same source image; aperture geometry and the source DR label remain unchanged. "
        "(C) Each regional crop rᵢ is encoded as zᵢ. Masked attention normalizes only across visible tokens before evidence aggregation h(m), producing the probability vector p(m) and predicted class ŷ(m). "
        "(D) The same fitted severity model produces p_full and p_limited under the two visibility conditions; paired features g are then used to obtain an exploratory limited-view underestimation-risk estimate q(g). "
        "Internal testing, one-region occlusion, and independent MMRDR-UWF evaluation form the evidence chain. All visibility restrictions are synthetic, and the paired-view audit requires predictions from both visibility conditions."
    )

    move_before_pair_details(image_paragraph, caption, [p3, equation, p4])
    style_tables_and_captions(doc)

    with tempfile.TemporaryDirectory(prefix="regionsuff_2_1_") as temp_dir:
        staged = Path(temp_dir) / args.output.name
        doc.save(staged)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(staged, args.output)

    print(f"Updated section 2.1 and Figure 1 caption in {args.output}")


if __name__ == "__main__":
    main()
