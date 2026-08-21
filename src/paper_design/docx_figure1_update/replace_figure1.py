from pathlib import Path

from docx import Document


SOURCE = Path("/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.4.docx")
FIGURE = Path(
    "/Users/tongyue/Documents/77f /7hao/paper_design/docx_figure1_update/figure1_2x/slide-1.png"
)
OUTPUT = Path("/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.5.docx")


def image_relationship_id(shape):
    blips = shape._inline.xpath(".//a:blip")
    if len(blips) != 1:
        raise RuntimeError("Expected exactly one embedded image relationship for Figure 1.")
    return blips[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")


def replace_caption(paragraph):
    body = (
        "RegionSuff study overview. (A) Technical gradability and task-relevant evidence "
        "availability are distinct: a clear image can omit evidence required for DR grading. "
        "(B) Region-level masks vary the visible 3×3 evidence tokens. Pixel-space apertures "
        "preserve a centered rectangle (f = 0.55), an area-matched circle (f = 0.55), or a "
        "severe rectangle (f = 0.35) from the same source image; retained content is magnified "
        "for display, while aperture geometry and the DR label are unchanged. (C) Each regional "
        "crop rᵢ is encoded as zᵢ; masked attention normalizes only over visible tokens before "
        "aggregation h(m), producing p(m) and ŷ(m). (D) The same fitted severity model yields "
        "p(m_full) and p(m_limited); paired features g then yield an exploratory underestimation-"
        "risk estimate q(g). Internal testing, one-region occlusion, and independent MMRDR-UWF "
        "replication form the evidence chain. Restrictions are synthetic; no clinician-adjudicated "
        "sufficiency labels or deployable single-view risk rule are claimed."
    )
    paragraph.clear()
    label = paragraph.add_run("Figure 1. ")
    label.bold = True
    paragraph.add_run(body)


def main():
    if not SOURCE.exists() or not FIGURE.exists():
        raise FileNotFoundError("The manuscript source or final Figure 1 PNG is missing.")

    document = Document(SOURCE)
    if len(document.inline_shapes) < 1:
        raise RuntimeError("No inline figures were found in the manuscript.")

    # Figure 1 is the inline drawing immediately before the Figure 1 caption.
    figure_shape = document.inline_shapes[0]
    figure_paragraph = document.paragraphs[21]
    caption = document.paragraphs[22]
    if not figure_paragraph._p.xpath(".//w:drawing") or not caption.text.startswith("Figure 1."):
        raise RuntimeError("The expected Figure 1 anchor/caption was not found; source was not modified.")

    relationship_id = image_relationship_id(figure_shape)
    image_part = document.part.related_parts[relationship_id]
    image_part._blob = FIGURE.read_bytes()
    image_part._image = None

    # Preserve the pre-existing figure dimensions and the local document layout.
    replace_caption(caption)

    paired_view_paragraph = document.paragraphs[20]
    paired_view_paragraph.text = (
        "The analyses were designed to characterize how the evidence available to the DR classifier "
        "changed as retinal coverage was restricted. The masks therefore represent controlled, synthetic "
        "restrictions of the observed retina rather than real pairs of conventional-field and UWF "
        "acquisitions. Because the datasets do not provide clinician-adjudicated labels of diagnostic "
        "sufficiency, the study focuses on model-specific evidence availability rather than clinical "
        "sufficiency itself. For the paired-view audit, the same fitted RegionSuff severity model was "
        "evaluated under a full-view mask m_full and a limited-view mask m_limited to obtain "
        "p(m_full) and p(m_limited); a separate downstream risk model then analyzed their paired "
        "differences. This module is an analytical comparison and does not operate on a single image."
    )

    document.core_properties.comments = "Figure 1 replaced with the final RegionSuff evidence-zoomed, uniform-LaTeX version."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
