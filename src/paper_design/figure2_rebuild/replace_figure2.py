from pathlib import Path

from docx import Document


SOURCE = Path("/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.5.docx")
FIGURE = Path("/Users/tongyue/Documents/77f /7hao/paper_design/figure2_rebuild/figure2_2x/slide-1.png")
OUTPUT = Path("/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.6.docx")


def image_relationship_id(shape):
    blips = shape._inline.xpath(".//a:blip")
    if len(blips) != 1:
        raise RuntimeError("Expected exactly one embedded image relationship for Figure 2.")
    return blips[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")


def replace_caption(paragraph):
    body = (
        "RegionSuff model and evaluation protocol. (A) The retinal foreground is divided into a "
        "3×3 evidence bank. Full, center, and cross masks are deterministic evaluation settings, "
        "whereas source training independently masks regional tokens with ρ = 0.30 and restores "
        "token 1 after an all-empty draw. (B) Every regional crop rᵢ passes through the shared "
        "frozen RetinaRadar encoder E and trainable projection P to produce zᵢ ∈ R¹²⁸. (C) Evidence "
        "scores are normalized only across visible tokens; their weighted aggregation yields p(m) "
        "and the primary DR severity prediction. Unweighted cross-entropy is the source-training "
        "loss. (D) The same severity model produces p_F = p(m_full) and p_L = p(m_limited); their "
        "18-dimensional paired feature g feeds the exploratory risk estimate q(g). The lower strip "
        "separates source fitting, internal testing, paired risk analysis, and independent MMRDR-UWF "
        "evaluation without target-domain training or selection."
    )
    paragraph.clear()
    label = paragraph.add_run("Figure 2. ")
    label.bold = True
    paragraph.add_run(body)


def main():
    if not SOURCE.exists() or not FIGURE.exists():
        raise FileNotFoundError("The manuscript source or final Figure 2 PNG is missing.")
    document = Document(SOURCE)
    if len(document.inline_shapes) < 2:
        raise RuntimeError("The expected Figure 2 inline image was not found.")

    figure_shape = document.inline_shapes[1]
    figure_paragraph = document.paragraphs[45]
    caption = document.paragraphs[46]
    if not figure_paragraph._p.xpath(".//w:drawing") or not caption.text.startswith("Figure 2."):
        raise RuntimeError("The expected Figure 2 anchor/caption was not found; source was not modified.")

    relationship_id = image_relationship_id(figure_shape)
    image_part = document.part.related_parts[relationship_id]
    image_part._blob = FIGURE.read_bytes()
    image_part._image = None
    replace_caption(caption)

    paired_view_paragraph = document.paragraphs[51]
    paired_view_paragraph.text = (
        "For each image, the same fitted RegionSuff severity model was evaluated twice: the full-view "
        "mask m_full produced p_F = p(m_full), and the center-plus-cross limited mask m_limited "
        "produced p_L = p(m_limited). These two outputs were passed to a downstream paired-view "
        "audit rather than treated as parallel outputs of the regional attention module. The resulting "
        "18-dimensional feature contains four probability blocks (12 values), three expected-severity "
        "terms, two predictive entropies, and one symmetric Kullback–Leibler divergence."
    )

    document.core_properties.comments = "Figure 2 replaced with the final RegionSuff algorithm and protocol figure."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
