from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.6.docx')
FIGURE = Path('/Users/tongyue/Documents/77f /7hao/paper_design/figure2_rebuild/mechanism-render-final/slide-1.png')
OUTPUT = Path('/Users/tongyue/Documents/ppf/第四论文/RegionSuff_1.7.docx')

CAPTION = (
    'Figure 2. RegionSuff mechanism: from visible retinal evidence to paired-view audit. '
    '(A) The retinal foreground is partitioned into a 3×3 evidence bank. The binary visibility map '
    'm specifies deterministic full, center, and center-plus-cross evaluation views. During source training, '
    'each token is independently masked with ρ = 0.30; an all-empty draw restores row-major token 0. '
    '(B) Each regional crop rᵢ is mapped by the shared frozen RetinaRadar EfficientNet-B0 encoder E and '
    'trainable projection P to a 128-dimensional token zᵢ. Attention is normalized only over visible tokens, '
    'which are aggregated to produce the three-class DR probability p(m) and severity prediction ŷ(m); '
    'unweighted cross-entropy is the source-training loss. (C) The same fitted severity model is evaluated '
    'under full and center-plus-cross masks to obtain p_F and p_L. Their 18-dimensional paired feature g '
    'is used by the exploratory paired-view risk estimate q(g), which requires both predictions and is not '
    'a deployable single-image triage rule.'
)


def replace_embedded_image(document_part, inline_shape, image_path: Path) -> None:
    """Replace the binary image behind an existing inline shape while preserving its Word geometry."""
    blip = inline_shape._inline.graphic.graphicData.pic.blipFill.blip
    relation_id = blip.get(qn('r:embed'))
    if not relation_id:
        raise RuntimeError('The Figure 2 inline image does not have an embedded image relationship.')
    image_part = document_part.related_parts[relation_id]
    image_part._blob = image_path.read_bytes()
    # python-docx memoizes decoded images; discard the cached object so the new PNG is used on save.
    if hasattr(image_part, '_image'):
        image_part._image = None


def set_paragraph_text(paragraph, value: str) -> None:
    paragraph.clear()
    paragraph.add_run(value)


def main() -> None:
    if not SOURCE.exists() or not FIGURE.exists():
        raise FileNotFoundError('Source manuscript or validated Figure 2 PNG was not found.')

    document = Document(SOURCE)
    if len(document.inline_shapes) < 2:
        raise RuntimeError(f'Expected at least two inline figures, found {len(document.inline_shapes)}.')

    # Figure 1 is inline shape 0; Figure 2 is shape 1 in the established manuscript layout.
    replace_embedded_image(document.part, document.inline_shapes[1], FIGURE)

    caption_index = None
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip().startswith('Figure 2.'):
            caption_index = index
            set_paragraph_text(paragraph, CAPTION)
            break
    if caption_index is None:
        raise RuntimeError('Could not locate the existing Figure 2 caption.')

    # Preserve the paper as a new numbered revision; never overwrite the source manuscript.
    document.save(OUTPUT)
    print({
        'source': str(SOURCE),
        'output': str(OUTPUT),
        'figure_shape_index': 1,
        'caption_paragraph_index': caption_index,
        'inline_shape_dimensions_emu': [
            document.inline_shapes[1].width,
            document.inline_shapes[1].height,
        ],
    })


if __name__ == '__main__':
    main()
